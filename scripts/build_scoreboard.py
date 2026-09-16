#!/usr/bin/env python3
"""Build The Called — The Scoreboard (.xlsx, imports cleanly into Google Sheets).

Run:  python3 scripts/build_scoreboard.py [--demo] [--logo PATH]
Out:  docs/The_Called_Scoreboard.xlsx  (or ..._DEMO.xlsx with sample data)

The logo is picked up automatically from docs/assets/the-called-logo.png when
that file exists; --logo overrides the path.

Design rules that matter (don't undo them without reading this):
  * No Sheets-only functions (QUERY/ARRAYFORMULA/SPARKLINE) — must open in both.
  * Conditional formatting NEVER references another tab: Google Sheets forbids it.
    Targets are mirrored onto the tab that needs them (Daily Log row 2, KPI col G).
  * Conditional-formatting ranges never overlap, so rule precedence can't surprise us.
  * Weekly Rollup is a rolling LAST 13 WEEKS — dashboard charts stay current with
    no empty tail.
  * Chart palette is validated for colour-blind separation (see PALETTE notes).
"""

import datetime
import os
import sys
import random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.text import RichText
from openpyxl.chart.marker import Marker
from openpyxl.drawing.text import (Paragraph, ParagraphProperties,
                                   CharacterProperties, RichTextProperties)
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment

from scoreboard_content import (INSTAGRAM, INTRO, MONEY, RATES, TRAPS, VOLUME)

# ---------------------------------------------------------------- brand
PALETTE = {
    # Sampled from The Called's logo — warm pewter/graphite on near-black.
    "dark":        "1B1916",  # logo background, lit corner — banners, table headers
    "ink":         "3A3833",  # letter shadow/recess — body headings
    "pewter":      "E1DFDB",  # letter highlight, brightest metal
    "tab_input":   "8E8A82",  # letter midtone — tabs the client types in

    # Chromatic accents. Brass is the natural partner to pewter; the four chart
    # hues below are validated together for colour-blind separation (worst
    # adjacent pair dE 11.2 deutan, normal-vision floor 21.4) — re-run
    # dataviz/scripts/validate_palette.js before changing any of them.
    "band":        "D4B579",  # soft sand — section bands, callouts, accents
    "band_soft":   "EFE1C6",  # tint of the same — today's row, target cells
    "brass":       "C8931A",  # chart series 1 ONLY — charts need the saturation
    "brass_light": "E3B341",  # brass that stays legible on the dark tiles
    "teal":        "0D8F7A",  # chart series 2
    "clay":        "C4552B",  # chart series 3
    "indigo":      "4A55A8",  # chart series 4

    "input":       "FCF3E3",  # warm parchment — "type here"
    "auto":        "F2F1EF",  # warm light grey — "leave this alone"
    "good":        "E7F0E5",
    "good_text":   "2F6B3A",
    "warn":        "FDF0D5",
    "warn_text":   "9A6B0F",
    "bad":         "F7E4E0",
    "bad_text":    "A83A2C",
    "line":        "D6D3CC",
    "muted":       "6E6A62",
}
FONT = "Arial"
CHART_SERIES = [PALETTE["brass"], PALETTE["teal"], PALETTE["clay"], PALETTE["indigo"]]

DAYS = 366          # one year; past that, copy the last row down or start a fresh copy
IG_WEEKS = 53
ROLL_WEEKS = 13

DL_FIRST, DL_LAST = 4, 4 + DAYS - 1
DATES = f"'Daily Log'!$A${DL_FIRST}:$A${DL_LAST}"
LOGGED = f"'Daily Log'!$Q${DL_FIRST}:$Q${DL_LAST}"

LOGO_DEFAULT = "docs/assets/the-called-logo.png"


def col(letter):
    return f"'Daily Log'!${letter}${DL_FIRST}:${letter}${DL_LAST}"


M = {"reels": "C", "stories": "D", "opener": "E", "follow": "F", "replies": "G",
     "pitched": "H", "booked": "I", "showed": "J", "closes": "K",
     "cash": "L", "revenue": "M", "followers": "N"}

# Setup input cells (column C), referenced everywhere
S_NAME, S_COACH, S_HANDLE, S_START, S_WEEK1 = "C5", "C6", "C7", "C8", "C9"
S_REELS, S_STORIES, S_OPENER, S_FOLLOW = "C12", "C13", "C14", "C15"
S_GOAL, S_PRICE = "C18", "C19"
STANDARD_CELLS = [S_REELS, S_STORIES, S_OPENER, S_FOLLOW]

# ---------------------------------------------------------------- helpers
thin = Side(style="thin", color=PALETTE["line"])
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
input_fill = PatternFill("solid", fgColor=PALETTE["input"])
auto_fill = PatternFill("solid", fgColor=PALETTE["auto"])
green = PatternFill("solid", bgColor=PALETTE["good"])
amber = PatternFill("solid", bgColor=PALETTE["warn"])
red = PatternFill("solid", bgColor=PALETTE["bad"])
brass_soft = PatternFill("solid", bgColor=PALETTE["band_soft"])

def banner(ws, text, last_col, row=1):
    """Dark banner across the top, with left padding that the logo drops into."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=PALETTE["dark"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=8)
    ws.row_dimensions[row].height = 46
    if LOGO_PATH:
        from openpyxl.drawing.image import Image
        img = Image(LOGO_PATH)
        img.height = 40
        img.width = 40
        img.anchor = f"A{row}"
        ws.add_image(img)


def subtitle(ws, text, last_col, row=2):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10, italic=True, color=PALETTE["muted"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 18


def section(ws, row, text, last_col):
    """Brass band with ink type — white on brass fails contrast, ink on brass sings."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10, bold=True, color=PALETTE["dark"])
    c.fill = PatternFill("solid", fgColor=PALETTE["band"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 21


def header_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PALETTE["dark"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 32


def label(ws, row, text, bold=False, indent=1, colnum=1):
    c = ws.cell(row=row, column=colnum, value=text)
    c.font = Font(name=FONT, size=10, bold=bold)
    c.alignment = Alignment(horizontal="left", indent=indent)
    return c


def widths(ws, mapping):
    for k, v in mapping.items():
        ws.column_dimensions[k].width = v


def tile(ws, row, c1, c2, label_text, formula, fmt="#,##0"):
    """Dark hero tile: small pewter label over a big brass number."""
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    lab = ws.cell(row=row, column=c1, value=label_text)
    lab.font = Font(name=FONT, size=8, bold=True, color=PALETTE["pewter"])
    lab.fill = PatternFill("solid", fgColor=PALETTE["dark"])
    lab.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 2, end_column=c2)
    val = ws.cell(row=row + 1, column=c1, value=formula)
    val.font = Font(name=FONT, size=20, bold=True, color=PALETTE["brass_light"])
    val.fill = PatternFill("solid", fgColor=PALETTE["dark"])
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = fmt
    ws.row_dimensions[row].height = 16
    ws.row_dimensions[row + 1].height = 16
    ws.row_dimensions[row + 2].height = 16


def sumifs(metric_col, start_ref, end_ref):
    return (f"SUMIFS({col(metric_col)},{DATES},\">=\"&{start_ref},"
            f"{DATES},\"<=\"&{end_ref})")


def streak(helper_col):
    """Current run length, tolerating 'hasn't logged today yet'."""
    rng = f"'Daily Log'!${helper_col}${DL_FIRST}:${helper_col}${DL_LAST}"
    today = f"INDEX({rng},MATCH(TODAY(),{DATES},0))"
    yday = f"INDEX({rng},MATCH(TODAY()-1,{DATES},0))"
    return f"=IFERROR(IF({today}>0,{today},{yday}),0)"


LOGO_PATH = None
if "--logo" in sys.argv:
    LOGO_PATH = sys.argv[sys.argv.index("--logo") + 1]
elif os.path.exists(LOGO_DEFAULT):
    LOGO_PATH = LOGO_DEFAULT

wb = Workbook()

# ================================================================ START HERE
ws = wb.active
ws.title = "Start Here"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 6, "B": 26, "C": 22, "D": 56, "E": 13, "F": 13, "G": 13, "H": 13})

banner(ws, "THE SCOREBOARD", 8)
subtitle(ws, "The work you put in, and what it turns into. Five minutes a week keeps it honest.", 8)

ws["B4"] = "You keep score, or you guess."
ws["B4"].font = Font(name=FONT, size=13, bold=True, color=PALETTE["ink"])
ws["B5"] = ("Everything you do lives in one sheet: posts, stories, DMs, replies, calls, cash. "
            "You fill in one row a day. The Scoreboard does the maths and tells you which part "
            "of your game is actually costing you money.")
ws.merge_cells("B5:H6")
ws["B5"].alignment = Alignment(wrap_text=True, vertical="top")
ws["B5"].font = Font(name=FONT, size=10)

section(ws, 8, "DO THESE IN ORDER", 8)
header_row(ws, 9, ["Step", "Tab", "When", "What you do"])
steps = [
    ("1", "Setup", "Once, right now", "Your name, your start date, your daily standards. Two minutes, never again."),
    ("2", "Daily Log", "Every day — 60 seconds", "Find today's row (it's highlighted) and fill in what you actually did."),
    ("3", "Instagram Tracker", "Every Sunday — 5 minutes", "Log the week's numbers off your Professional Dashboard."),
    ("4", "This Week", "Every Monday — 2 minutes", "What you hit, what you missed, and the one thing to fix this week."),
    ("5", "KPI Dashboard", "Before every coaching call", "The full picture: rates, trends, and where the money's leaking."),
]
r = 10
for n, tab, when, what in steps:
    c = ws.cell(row=r, column=1, value=n)
    c.font = Font(name=FONT, size=14, bold=True, color=PALETTE["ink"])
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=2, value=tab).font = Font(name=FONT, size=11, bold=True, color=PALETTE["ink"])
    ws.cell(row=r, column=3, value=when).font = Font(name=FONT, size=9, color=PALETTE["muted"])
    ws.cell(row=r, column=4, value=what).font = Font(name=FONT, size=10)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    for cc in (2, 3, 4):
        ws.cell(row=r, column=cc).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=(cc == 4), indent=1)
    for cc in range(1, 9):
        ws.cell(row=r, column=cc).border = BOX
        if r % 2 == 0:
            ws.cell(row=r, column=cc).fill = auto_fill
    ws.row_dimensions[r].height = 30
    r += 1

section(ws, 16, "REFERENCE — open these when you need them", 8)
for i, (nm, desc) in enumerate((
        ("Cheat Sheet", "What every number means, what a low one is telling you, and exactly what to "
                        "do about it. Read it once, then whenever a number looks wrong."),
        ("Weekly Rollup", "Feeds the charts on the dashboard. Nothing to fill in. Leave it alone."))):
    rr = 17 + i
    ws.cell(row=rr, column=2, value=nm).font = Font(name=FONT, size=10, bold=True)
    d = ws.cell(row=rr, column=4, value=desc)
    d.font = Font(name=FONT, size=10, color=PALETTE["muted"])
    d.alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=8)
    ws.row_dimensions[rr].height = 26

section(ws, 19, "THE RULES — read once, saves you a headache", 8)
rules = [
    ("Cream cells are yours.", "Grey cells calculate themselves. If it's grey, don't type in it."),
    ("Log a zero as 0, never blank.", "Blank means 'didn't track'. Blanks make your averages flatter you."),
    ("Don't delete columns or rows.", "It breaks the maths. Don't need a column? Right-click and Hide it."),
    ("Backfill a missed day.", "A gap in the log is a lie in your numbers. Fill it in when you remember."),
]
r = 20
for head, body in rules:
    ws.cell(row=r, column=2, value=head).font = Font(name=FONT, size=10, bold=True, color=PALETTE["ink"])
    ws.cell(row=r, column=4, value=body).font = Font(name=FONT, size=10)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.row_dimensions[r] = ws.row_dimensions[r]
    ws.row_dimensions[r].height = 22
    r += 1

section(ws, 25, "EXAMPLE — what one good day looks like", 8)
ex_head = ["Date", "Reels", "Stories", "Opener DMs", "Follow-Ups", "Replies", "Pitched", "Booked"]
ex_vals = ["Mon 15 Sep", 5, 5, 20, 10, 4, 2, 1]
for i, (h, v) in enumerate(zip(ex_head, ex_vals)):
    hc = ws.cell(row=26, column=1 + i, value=h)
    hc.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
    hc.fill = PatternFill("solid", fgColor=PALETTE["dark"])
    hc.alignment = Alignment(horizontal="center")
    vc = ws.cell(row=27, column=1 + i, value=v)
    vc.font = Font(name=FONT, size=11, bold=True, color=PALETTE["good_text"])
    vc.fill = PatternFill("solid", fgColor=PALETTE["good"])
    vc.alignment = Alignment(horizontal="center")
    vc.border = BOX
ws["A28"] = "All four standards hit → the day scores 4/4, and your perfect-day streak goes up."
ws["A28"].font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])
ws.merge_cells("A28:H28")

section(ws, 30, "WALKTHROUGH VIDEOS", 8)
for i, name in enumerate(["Setup + Daily Log", "KPI Dashboard + This Week",
                          "Instagram Tracker", "Content Planner (Notion)"]):
    r = 31 + i
    ws.cell(row=r, column=2, value=name).font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r, column=4, value="[paste Loom link]")
    c.font = Font(name=FONT, size=10, color=PALETTE["indigo"])
    c.fill = input_fill
    c.border = BOX
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)

# ================================================================ SETUP
st = wb.create_sheet("Setup")
st.sheet_view.showGridLines = False
widths(st, {"A": 4, "B": 32, "C": 18, "D": 54})
banner(st, "SETUP — two minutes, once", 4)
subtitle(st, "Fill in the cream cells. Every other tab reads from here.", 4)


def setup_row(row, lbl, value, note, fmt=None, is_input=True):
    label(st, row, lbl, bold=True, colnum=2)
    c = st.cell(row=row, column=3, value=value)
    c.font = Font(name=FONT, size=11, bold=is_input)
    c.fill = input_fill if is_input else auto_fill
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
    if fmt:
        c.number_format = fmt
    n = st.cell(row=row, column=4, value=note)
    n.font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])
    st.row_dimensions[row].height = 20
    return c


section(st, 4, "WHO THIS IS", 4)
setup_row(5, "Client name", "", "Shows on your dashboard.")
setup_row(6, "Coach / CSM", "", "Who you report to inside The Called.")
setup_row(7, "Instagram handle", "", "e.g. @yourhandle")
setup_row(8, "Tracking start date", "=TODAY()", "Day 1. This sets every date in the Daily Log.",
          fmt="ddd d mmm yyyy")
setup_row(9, "First tracking week (auto)", "=C8-WEEKDAY(C8,2)+1",
          "Monday of your start week. Calculated — don't edit.",
          fmt="ddd d mmm yyyy", is_input=False)

section(st, 11, "YOUR DAILY STANDARDS — the floor, not the goal", 4)
setup_row(12, "Reels / posts per day", 5, "Default 5.")
setup_row(13, "Stories per day", 5, "Default 5.")
setup_row(14, "Opener DMs per day", 20, "Default 20.")
setup_row(15, "Follow-up DMs per day", 10, "Default 10.")
st["D12"].comment = Comment(
    "The Called's standard: 5 posts, 5 stories, 20 opener DMs, 10 follow-ups. "
    "Change them per client if their programme says something different — every tab "
    "re-grades itself from these four cells.", "The Called")

section(st, 17, "WHAT YOU'RE PLAYING FOR", 4)
setup_row(18, "Monthly cash collected goal", 10000, "Drives the goal tile on the dashboard.",
          fmt='"$"#,##0')
setup_row(19, "Your offer price", 0, "Optional — for your own maths.", fmt='"$"#,##0')

# ================================================================ DAILY LOG
dl = wb.create_sheet("Daily Log")
dl.sheet_view.showGridLines = False
banner(dl, "DAILY LOG — one row per day", 20)

lab = dl.cell(row=2, column=1, value="Your standard →")
lab.font = Font(name=FONT, size=9, bold=True, color=PALETTE["dark"])
lab.alignment = Alignment(horizontal="right")
lab.fill = PatternFill("solid", fgColor=PALETTE["band"])
dl["B2"].fill = PatternFill("solid", fgColor=PALETTE["band"])
for c_, ref in (("C", S_REELS), ("D", S_STORIES), ("E", S_OPENER), ("F", S_FOLLOW)):
    cell = dl[f"{c_}2"]
    cell.value = f"=Setup!${ref[0]}${ref[1:]}"
    cell.font = Font(name=FONT, size=10, bold=True, color=PALETTE["dark"])
    cell.fill = PatternFill("solid", fgColor=PALETTE["band"])
    cell.alignment = Alignment(horizontal="center")
dl.merge_cells("H2:P2")
note = dl["H2"]
note.value = "Today's row is highlighted. Log a zero as 0 — blank means you didn't track."
note.font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])
note.alignment = Alignment(horizontal="left", vertical="center", indent=1)
dl.row_dimensions[2].height = 19

dl_headers = ["Date", "Day", "Reels /\nPosts", "Stories", "Opener\nDMs Sent",
              "Follow-Up\nDMs Sent", "Replies", "Calls\nPitched", "Calls\nBooked",
              "Calls\nShowed", "Closes", "Cash\nCollected", "Revenue\nGenerated",
              "Followers\n(end of day)", "Standards\nHit (of 4)", "Notes",
              "Logged?\n(auto)", "Day\nStreak", "Perfect\nStreak", "Logged on\n(auto)"]
header_row(dl, 3, dl_headers)
widths(dl, {"A": 13, "B": 6, "C": 8, "D": 8, "E": 9, "F": 10, "G": 9, "H": 9, "I": 9,
            "J": 9, "K": 8, "L": 11, "M": 11, "N": 11, "O": 10, "P": 32, "Q": 8,
            "R": 8, "S": 8, "T": 12})
dl.freeze_panes = "C4"

money_cols = {"L", "M"}
input_cols = list("CDEFGHIJKLMN")
for i in range(DAYS):
    r = DL_FIRST + i
    a = dl.cell(row=r, column=1, value=f"=Setup!${S_START[0]}${S_START[1:]}" if i == 0
                else f"=A{r-1}+1")
    a.number_format = "ddd d mmm"
    a.font = Font(name=FONT, size=10)
    a.fill = auto_fill
    b = dl.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"ddd"))')
    b.font = Font(name=FONT, size=9, color=PALETTE["muted"])
    b.fill = auto_fill
    b.alignment = Alignment(horizontal="center")
    for letter in input_cols:
        c = dl[f"{letter}{r}"]
        c.fill = input_fill
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        c.number_format = '"$"#,##0' if letter in money_cols else "#,##0"
    o = dl.cell(row=r, column=15,
                value=(f'=IF($Q{r}=0,"",IF(C{r}>=C$2,1,0)+IF(D{r}>=D$2,1,0)'
                       f'+IF(E{r}>=E$2,1,0)+IF(F{r}>=F$2,1,0))'))
    o.alignment = Alignment(horizontal="center")
    o.font = Font(name=FONT, size=10, bold=True)
    o.border = BOX
    p = dl.cell(row=r, column=16)
    p.fill = input_fill
    p.font = Font(name=FONT, size=9)
    p.border = BOX
    dl.cell(row=r, column=17, value=f'=IF(COUNT(C{r}:M{r})>0,1,0)')
    dl.cell(row=r, column=18,
            value=f'=IF(Q{r}=0,0,1)' if i == 0 else f'=IF(Q{r}=0,0,R{r-1}+1)')
    dl.cell(row=r, column=19,
            value=(f'=IF(O{r}="",0,IF(O{r}=4,1,0))' if i == 0
                   else f'=IF(O{r}="",0,IF(O{r}=4,S{r-1}+1,0))'))
    t = dl.cell(row=r, column=20, value=f'=IF(Q{r}=1,A{r},"")')
    t.number_format = "d mmm"
    for cidx in (17, 18, 19, 20):
        c = dl.cell(row=r, column=cidx)
        c.fill = auto_fill
        c.font = Font(name=FONT, size=8, color=PALETTE["muted"])
        c.alignment = Alignment(horizontal="center")

dl["Q3"].comment = Comment(
    "Three auto columns. Logged? is 1 on days you filled something in. Day Streak counts "
    "days logged in a row. Perfect Streak counts 4/4 days in a row. The dashboard reads "
    "all three — don't type in them or delete them.", "The Called")

dv_int = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Whole numbers only",
                        error="Enter a whole number of 0 or more.")
dv_money = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                          allow_blank=True, showErrorMessage=True,
                          errorTitle="Amount only",
                          error="Enter a dollar amount of 0 or more.")
dl.add_data_validation(dv_int)
dl.add_data_validation(dv_money)
for letter in "CDEFGHIJKN":
    dv_int.add(f"{letter}{DL_FIRST}:{letter}{DL_LAST}")
for letter in "LM":
    dv_money.add(f"{letter}{DL_FIRST}:{letter}{DL_LAST}")

# CF ranges below never overlap, so rule precedence can't surprise anyone.
rng = f"C{DL_FIRST}:F{DL_LAST}"
dl.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($Q{DL_FIRST}=1,C{DL_FIRST}>=C$2)'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"]), stopIfTrue=False))
dl.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($Q{DL_FIRST}=1,C{DL_FIRST}<C$2)'], fill=red,
    font=Font(name=FONT, color=PALETTE["bad_text"]), stopIfTrue=False))
dl.conditional_formatting.add(f"O{DL_FIRST}:O{DL_LAST}", FormulaRule(
    formula=[f'$O{DL_FIRST}=4'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"]), stopIfTrue=False))
dl.conditional_formatting.add(f"O{DL_FIRST}:O{DL_LAST}", FormulaRule(
    formula=[f'AND($Q{DL_FIRST}=1,$O{DL_FIRST}<2)'], fill=red,
    font=Font(name=FONT, bold=True, color=PALETTE["bad_text"]), stopIfTrue=False))
for zone in (f"A{DL_FIRST}:B{DL_LAST}", f"G{DL_FIRST}:N{DL_LAST}",
             f"P{DL_FIRST}:T{DL_LAST}"):
    dl.conditional_formatting.add(zone, FormulaRule(
        formula=[f'$A{DL_FIRST}=TODAY()'], fill=brass_soft,
        font=Font(name=FONT, bold=True, color=PALETTE["dark"]), stopIfTrue=False))

# ================================================================ THIS WEEK
tw = wb.create_sheet("This Week")
tw.sheet_view.showGridLines = False
widths(tw, {"A": 30, "B": 15, "C": 18, "D": 14, "E": 14, "F": 12, "G": 18})
banner(tw, "THIS WEEK", 7)
tw["A2"] = ('="Week of "&TEXT(B3,"d mmm")&" — day "&D3&" of 7.   Last week: "'
            '&TEXT(F3,"d mmm")&" to "&TEXT(G3,"d mmm")')
tw.merge_cells("A2:G2")
tw["A2"].font = Font(name=FONT, size=10, italic=True, color=PALETTE["muted"])
tw["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

helpers = [("A3", "Week starts"), ("B3", "=TODAY()-WEEKDAY(TODAY(),2)+1"),
           ("C3", "Day of week"), ("D3", "=TODAY()-B3+1"),
           ("E3", "Last week"), ("F3", "=B3-7"), ("G3", "=B3-1")]
for ref, val in helpers:
    c = tw[ref]
    c.value = val
    c.font = Font(name=FONT, size=8, italic=True, color=PALETTE["muted"])
    c.alignment = Alignment(horizontal="center")
    if ref in ("B3", "F3", "G3"):
        c.number_format = "d mmm"

c1 = tw["A5"]
c1.value = ('=IF(SUM(B10:B13)=0,'
            '"Log a couple of days and this line tells you what to fix first.",'
            'IF(MIN($F$10:$F$13)>=1,"Every standard is at or ahead of pace. Hold it — and push volume.",'
            '"Fix this first: "&INDEX($A$10:$A$13,MATCH(MIN($F$10:$F$13),$F$10:$F$13,0))'
            '&" — you are at "&TEXT(MIN($F$10:$F$13),"0%")&" of your target for end of today."))')
tw.merge_cells("A5:G5")
c1.font = Font(name=FONT, size=12, bold=True, color=PALETTE["dark"])
c1.fill = PatternFill("solid", fgColor=PALETTE["band"])
c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
tw.row_dimensions[5].height = 28

c2 = tw["A6"]
c2.value = ('=IF(COUNT(\'Weekly Rollup\'!$C$4:$C$16)<3,'
            '"Three weeks in, this line names the weakest step in your funnel.",'
            'IF(COUNT($D$27:$D$31)=0,"Log some calls and this line compares your rates to your average.",'
            'IF(MIN($D$27:$D$31)>=0,"Every rate is at or above your average this week — so push volume.",'
            '"Biggest drop vs your average: "&INDEX($A$27:$A$31,MATCH(MIN($D$27:$D$31),$D$27:$D$31,0))'
            '&", "&TEXT(ABS(MIN($D$27:$D$31)),"0.0%")&" below your usual.")))')
tw.merge_cells("A6:G6")
c2.font = Font(name=FONT, size=11, color=PALETTE["dark"])
c2.fill = PatternFill("solid", fgColor=PALETTE["band_soft"])
c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
tw.row_dimensions[6].height = 24


def tw_cell(ref, value, fmt=None, bold=False, fill=None, center=True):
    c = tw[ref]
    c.value = value
    c.font = Font(name=FONT, size=10, bold=bold)
    c.border = BOX
    c.fill = fill or auto_fill
    c.alignment = Alignment(horizontal="center" if center else "left", indent=0 if center else 1)
    if fmt:
        c.number_format = fmt
    return c


section(tw, 8, "THE WORK — this part you control", 7)
header_row(tw, 9, ["Standard", "This week so far", "Target by\nend of today",
                   "Ahead / behind", "Weekly target", "% of pace",
                   "Next week: aim per day"])
work = [("Reels / posts", "reels", S_REELS), ("Stories", "stories", S_STORIES),
        ("Opener DMs", "opener", S_OPENER), ("Follow-up DMs", "follow", S_FOLLOW)]
for i, (name, key, sref) in enumerate(work):
    r = 10 + i
    s = f"Setup!${sref[0]}${sref[1:]}"
    label(tw, r, name, bold=True)
    tw[f"A{r}"].border = BOX
    tw_cell(f"B{r}", "=" + sumifs(M[key], "$B$3", "TODAY()"), "#,##0", bold=True)
    tw_cell(f"C{r}", f"={s}*$D$3", "#,##0")
    tw_cell(f"D{r}", f"=B{r}-C{r}", "+#,##0;-#,##0;0", bold=True)
    tw_cell(f"E{r}", f"={s}*7", "#,##0")
    tw_cell(f"F{r}", f"=IFERROR(B{r}/C{r},0)", "0%", bold=True)
    tw_cell(f"G{r}", f"=IF(F{r}>=1,ROUND({s}*1.1,0),{s})", "#,##0",
            fill=PatternFill("solid", fgColor=PALETTE["band_soft"]), bold=True)

section(tw, 15, "THE RESULT — this part follows", 7)
header_row(tw, 16, ["Metric", "This week so far", "Last week", "Change"])
result = [("Total DMs sent", None), ("Replies", "replies"), ("Calls pitched", "pitched"),
          ("Calls booked", "booked"), ("Calls showed", "showed"), ("Closes", "closes"),
          ("Cash collected", "cash")]
for i, (name, key) in enumerate(result):
    r = 17 + i
    label(tw, r, name, bold=(key in (None, "cash")))
    tw[f"A{r}"].border = BOX
    fmt = '"$"#,##0' if key == "cash" else "#,##0"
    if key is None:
        this_w = "=" + sumifs(M["opener"], "$B$3", "TODAY()") + "+" + sumifs(M["follow"], "$B$3", "TODAY()")
        last_w = "=" + sumifs(M["opener"], "$F$3", "$G$3") + "+" + sumifs(M["follow"], "$F$3", "$G$3")
    else:
        this_w = "=" + sumifs(M[key], "$B$3", "TODAY()")
        last_w = "=" + sumifs(M[key], "$F$3", "$G$3")
    tw_cell(f"B{r}", this_w, fmt, bold=True)
    tw_cell(f"C{r}", last_w, fmt)
    tw_cell(f"D{r}", f"=B{r}-C{r}", ('"+$"#,##0;"-$"#,##0;0' if key == "cash"
                                     else "+#,##0;-#,##0;0"), bold=True)

section(tw, 25, "YOUR RATES vs YOUR OWN 13-WEEK AVERAGE", 7)
header_row(tw, 26, ["Rate", "This week", "Your average", "Difference"])
rates = [("Reply rate", "B18/B17", "P"), ("Pitch rate", "B19/B18", "Q"),
         ("Book rate", "B20/B19", "R"), ("Show rate", "B21/B20", "S"),
         ("Close rate", "B22/B21", "T")]
for i, (name, expr, wcol) in enumerate(rates):
    r = 27 + i
    label(tw, r, name, bold=True)
    tw[f"A{r}"].border = BOX
    tw_cell(f"B{r}", f'=IFERROR({expr},"")', "0.0%", bold=True)
    tw_cell(f"C{r}", f'=IFERROR(AVERAGE(\'Weekly Rollup\'!${wcol}$4:${wcol}$16),"")', "0.0%")
    tw_cell(f"D{r}", f'=IFERROR(B{r}-C{r},"")', "+0.0%;-0.0%;0.0%", bold=True)

section(tw, 33, "AIM FOR NEXT WEEK", 7)
head = tw["A34"]
head.value = ('=IF(SUM(B10:B13)=0,"Log this week first — next week\'s aim builds off it.",'
              '"Next week, the one to protect: "'
              '&INDEX($A$10:$A$13,MATCH(MIN($F$10:$F$13),$F$10:$F$13,0))&" — "'
              '&INDEX($G$10:$G$13,MATCH(MIN($F$10:$F$13),$F$10:$F$13,0))&" a day, every day.")')
tw.merge_cells("A34:G34")
head.font = Font(name=FONT, size=12, bold=True, color=PALETTE["dark"])
head.fill = PatternFill("solid", fgColor=PALETTE["band_soft"])
head.alignment = Alignment(horizontal="left", vertical="center", indent=1)
tw.row_dimensions[34].height = 26

header_row(tw, 35, ["Standard", "This week", "Per day", "Week total", "Why"])
tw.merge_cells("E35:G35")
for i, (name, key, sref) in enumerate(work):
    r, src = 36 + i, 10 + i
    label(tw, r, name, bold=True)
    tw[f"A{r}"].border = BOX
    tw_cell(f"B{r}", f"=B{src}", "#,##0")
    tw_cell(f"C{r}", f"=G{src}", "#,##0",
            fill=PatternFill("solid", fgColor=PALETTE["band_soft"]), bold=True)
    tw_cell(f"D{r}", f"=C{r}*7", "#,##0", bold=True)
    tw.merge_cells(f"E{r}:G{r}")
    w = tw_cell(f"E{r}", f'=IF(F{src}>=1,"Held the line — push it 10%.",'
                         f'"Missed it. Hold the standard and close the gap.")',
                center=False)
    w.font = Font(name=FONT, size=10, italic=True)

tw["A41"] = ("Per day nudges up 10% on anything you hit this week and holds the standard on "
             "anything you didn't. It's a suggestion — your actual standards live on Setup.")
tw.merge_cells("A41:G41")
tw["A41"].font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])
tw["A41"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
tw.row_dimensions[41].height = 26

for zone, good_f, bad_f in ((f"D10:D13", "D10>=0", "D10<0"),
                            (f"D17:D23", "D17>0", "D17<0"),
                            (f"D27:D31", 'AND(D27<>"",D27>0)', 'AND(D27<>"",D27<0)')):
    tw.conditional_formatting.add(zone, FormulaRule(
        formula=[good_f], fill=green, font=Font(name=FONT, bold=True, color=PALETTE["good_text"])))
    tw.conditional_formatting.add(zone, FormulaRule(
        formula=[bad_f], fill=red, font=Font(name=FONT, bold=True, color=PALETTE["bad_text"])))
tw.conditional_formatting.add("F10:F13", FormulaRule(
    formula=["F10>=1"], fill=green, font=Font(name=FONT, bold=True, color=PALETTE["good_text"])))
tw.conditional_formatting.add("F10:F13", FormulaRule(
    formula=["AND(F10>=0.8,F10<1)"], fill=amber,
    font=Font(name=FONT, bold=True, color=PALETTE["warn_text"])))
tw.conditional_formatting.add("F10:F13", FormulaRule(
    formula=["F10<0.8"], fill=red, font=Font(name=FONT, bold=True, color=PALETTE["bad_text"])))

# ================================================================ WEEKLY ROLLUP
wr = wb.create_sheet("Weekly Rollup")
wr.sheet_view.showGridLines = False
WR_FIRST = 4
WR_LAST = WR_FIRST + ROLL_WEEKS - 1
banner(wr, "WEEKLY ROLLUP — your last 13 weeks", 20)
subtitle(wr, "All automatic. The dashboard charts read from here. Nothing to fill in.", 20)

header_row(wr, 3, ["Week Starting", "Week", "Days\nLogged", "Reels /\nPosts", "Stories",
                   "Opener\nDMs", "Follow-Up\nDMs", "Total DMs\nSent", "Replies",
                   "Calls\nPitched", "Calls\nBooked", "Calls\nShowed", "Closes",
                   "Cash\nCollected", "Revenue", "Reply\nRate", "Pitch\nRate",
                   "Book\nRate", "Show\nRate", "Close\nRate"])
widths(wr, {"A": 14, "B": 9, "C": 8, "D": 8, "E": 8, "F": 9, "G": 9, "H": 9, "I": 9,
            "J": 9, "K": 9, "L": 9, "M": 8, "N": 11, "O": 11, "P": 8, "Q": 8, "R": 8,
            "S": 8, "T": 8})
wr.freeze_panes = "C4"

GUARD = 'IF($A{r}<Setup!${wk}${wkn},"",{body})'
for i in range(ROLL_WEEKS):
    r = WR_FIRST + i
    first = f"=TODAY()-WEEKDAY(TODAY(),2)+1-{7 * (ROLL_WEEKS - 1)}"
    a = wr.cell(row=r, column=1, value=first if i == 0 else f"=A{r-1}+7")
    a.number_format = "ddd d mmm"
    a.font = Font(name=FONT, size=10)
    b = wr.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"d mmm"))')
    b.font = Font(name=FONT, size=9, color=PALETTE["muted"])
    b.alignment = Alignment(horizontal="center")

    def guarded(body):
        return "=" + GUARD.format(r=r, wk=S_WEEK1[0], wkn=S_WEEK1[1:], body=body)

    wr.cell(row=r, column=3, value=guarded(
        f'SUMIFS({LOGGED},{DATES},">="&$A{r},{DATES},"<="&$A{r}+6)'))
    for cidx, key in ((4, "reels"), (5, "stories"), (6, "opener"), (7, "follow"),
                      (9, "replies"), (10, "pitched"), (11, "booked"), (12, "showed"),
                      (13, "closes"), (14, "cash"), (15, "revenue")):
        wr.cell(row=r, column=cidx, value=guarded(
            f'SUMIFS({col(M[key])},{DATES},">="&$A{r},{DATES},"<="&$A{r}+6)'))
    wr.cell(row=r, column=8, value=guarded(f'IFERROR(F{r}+G{r},"")'))
    for cidx, expr in ((16, f"I{r}/H{r}"), (17, f"J{r}/I{r}"), (18, f"K{r}/J{r}"),
                       (19, f"L{r}/K{r}"), (20, f"M{r}/L{r}")):
        wr.cell(row=r, column=cidx, value=f'=IFERROR({expr},"")')
    for cidx in range(3, 21):
        c = wr.cell(row=r, column=cidx)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        c.fill = auto_fill
        c.number_format = ('"$"#,##0' if cidx in (14, 15)
                           else "0.0%" if cidx >= 16 else "#,##0")

# ================================================================ KPI DASHBOARD
kd = wb.create_sheet("KPI Dashboard")
kd.sheet_view.showGridLines = False
banner(kd, "KPI DASHBOARD", 11)
subtitle(kd, "100% automatic. If a number looks wrong, the fix is in the Daily Log — not here.", 11)
widths(kd, {"A": 32, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13, "G": 11,
            "H": 13, "I": 13, "J": 13, "K": 13})

lbl = kd["A5"]
lbl.value = "AT A GLANCE"
lbl.font = Font(name=FONT, size=10, bold=True, color=PALETTE["ink"])
lbl.alignment = Alignment(horizontal="left", vertical="center", indent=1)

tile(kd, 4, 2, 3, "CASH THIS MONTH", "=" + sumifs(M["cash"], "$D$10", "$D$11"), '"$"#,##0')
tile(kd, 4, 4, 5, "% OF MONTHLY GOAL", f"=IFERROR(B5/Setup!${S_GOAL[0]}${S_GOAL[1:]},0)", "0%")
tile(kd, 4, 6, 7, "CALLS BOOKED THIS WEEK", "=" + sumifs(M["booked"], "$B$10", "$B$11"))
tile(kd, 4, 8, 9, "DAYS LOGGED IN A ROW", streak("R"))
tile(kd, 4, 10, 11, "PERFECT DAYS IN A ROW", streak("S"))

# CSM check: is this sheet current? Read before a call, not chased by the sheet.
g7 = kd["G7"]
g7.value = ('=IF(MAX(\'Daily Log\'!$T$4:$T$369)=0,-1,'
            'TODAY()-MAX(\'Daily Log\'!$T$4:$T$369))')
g7.font = Font(name=FONT, size=8, color=PALETTE["auto"])
g7.alignment = Alignment(horizontal="center")
kd.merge_cells("A7:E7")
le = kd["A7"]
le.value = ('="Last entry: "&IF($G$7=-1,"nothing logged yet",'
            'IF($G$7=0,"today",IF($G$7=1,"yesterday",$G$7&" days ago")))')
le.font = Font(name=FONT, size=10, bold=True)
le.alignment = Alignment(horizontal="left", vertical="center", indent=1)
kd.row_dimensions[7].height = 19
kd.conditional_formatting.add("A7:E7", FormulaRule(
    formula=["$G$7>3"], fill=red, font=Font(name=FONT, bold=True, color=PALETTE["bad_text"])))
kd.conditional_formatting.add("A7:E7", FormulaRule(
    formula=["AND($G$7>=0,$G$7<=1)"], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"])))

periods = [("B", "This Week", "=TODAY()-WEEKDAY(TODAY(),2)+1", "=TODAY()"),
           ("C", "Last Week", "=B10-7", "=B10-1"),
           ("D", "This Month", "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "=TODAY()"),
           ("E", "Last Month", "=EOMONTH(TODAY(),-2)+1", "=EOMONTH(TODAY(),-1)"),
           ("F", "All Time", f"=Setup!${S_START[0]}${S_START[1:]}", "=TODAY()")]
for letter, name, s, e in periods:
    h = kd[f"{letter}8"]
    h.value = name
    h.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=PALETTE["ink"])
    h.alignment = Alignment(horizontal="center", vertical="center")
    lb = kd[f"{letter}9"]
    lb.value = f'=TEXT({letter}10,"d mmm")&" – "&TEXT({letter}11,"d mmm")'
    lb.font = Font(name=FONT, size=8, color=PALETTE["muted"])
    lb.alignment = Alignment(horizontal="center")
    for row, f in ((10, s), (11, e)):
        c = kd[f"{letter}{row}"]
        c.value = f
        c.number_format = "d mmm yyyy"
        c.font = Font(name=FONT, size=8, color=PALETTE["muted"])
        c.alignment = Alignment(horizontal="center")
kd.row_dimensions[8].height = 20
for row, txt in ((10, "Period start"), (11, "Period end")):
    c = label(kd, row, txt)
    c.font = Font(name=FONT, size=8, italic=True, color=PALETTE["muted"])


def metric_row(row, lbl_text, body_fn, fmt="#,##0", bold=False, indent=2):
    label(kd, row, lbl_text, bold=bold, indent=indent)
    for letter, *_ in periods:
        c = kd[f"{letter}{row}"]
        c.value = "=" + body_fn(letter)
        c.number_format = fmt
        c.font = Font(name=FONT, size=10, bold=bold)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        c.fill = auto_fill


section(kd, 13, "ACTIVITY — what you put in", 11)
metric_row(14, "Days Logged",
           lambda L: f'SUMIFS({LOGGED},{DATES},">="&{L}$10,{DATES},"<="&{L}$11)')
for row, lbl_text, key in ((15, "Reels / Posts", "reels"), (16, "Stories", "stories"),
                           (17, "Opener DMs Sent", "opener"), (18, "Follow-Up DMs Sent", "follow")):
    metric_row(row, lbl_text, lambda L, k=key: sumifs(M[k], f"{L}$10", f"{L}$11"))
metric_row(19, "Total DMs Sent", lambda L: f'{L}17+{L}18', bold=True)
metric_row(20, "Replies", lambda L: sumifs(M["replies"], f"{L}$10", f"{L}$11"))

section(kd, 22, "SALES — what it turned into", 11)
for row, lbl_text, key in ((23, "Calls Pitched", "pitched"), (24, "Calls Booked", "booked"),
                           (25, "Calls Showed", "showed"), (26, "Closes", "closes")):
    metric_row(row, lbl_text, lambda L, k=key: sumifs(M[k], f"{L}$10", f"{L}$11"))
metric_row(27, "Cash Collected", lambda L: sumifs(M["cash"], f"{L}$10", f"{L}$11"),
           fmt='"$"#,##0', bold=True)
metric_row(28, "Revenue Generated", lambda L: sumifs(M["revenue"], f"{L}$10", f"{L}$11"),
           fmt='"$"#,##0')

# band stops at F so the healthy-floor column keeps its own header cell
section(kd, 30, "CONVERSION RATES — where you're actually leaking", 6)
for cell_ref, text in (("G30", "Red\nbelow"), ("H30", "The bar")):
    hf = kd[cell_ref]
    hf.value = text
    hf.font = Font(name=FONT, size=8, bold=True, color="FFFFFF")
    hf.fill = PatternFill("solid", fgColor=PALETTE["dark"])
    hf.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
rate_rows = [(31, "Reply Rate (replies ÷ DMs sent)", "20", "19", "0.0%"),
             (32, "Pitch Rate (pitched ÷ replies)", "23", "20", "0.0%"),
             (33, "Book Rate (booked ÷ pitched)", "24", "23", "0.0%"),
             (34, "Show Rate (showed ÷ booked)", "25", "24", "0.0%"),
             (35, "Close Rate (closes ÷ showed)", "26", "25", "0.0%"),
             (36, "DMs per Call Booked", "19", "24", "0.0"),
             (37, "Cash per Call Booked", "27", "24", '"$"#,##0')]
for row, lbl_text, num, den, fmt in rate_rows:
    metric_row(row, lbl_text, lambda L, n=num, d=den: f'IFERROR({L}{n}/{L}{d},"")', fmt=fmt)
# floors mirrored onto this tab: Google Sheets conditional formatting cannot
# reach across tabs, and these are the numbers the colour rules compare against
# both thresholds mirrored onto this tab: Google Sheets conditional formatting
# cannot reach across tabs, and these are what the colour rules compare against
for i, row in enumerate(range(31, 36)):
    for col_letter, src_col in (("G", "C"), ("H", "D")):   # C = red below, D = the bar
        g = kd[f"{col_letter}{row}"]
        g.value = f"='Cheat Sheet'!${src_col}${11 + i}"
        g.number_format = "0%"
        g.font = Font(name=FONT, size=10, bold=True, color=PALETTE["ink"])
        g.alignment = Alignment(horizontal="center")
        g.fill = PatternFill("solid", fgColor=PALETTE["band_soft"])
        g.border = BOX
# three tiers: at/above the bar is green, genuinely broken is red, and the wide
# middle is amber — a client working well but short of the bar is not a failure
kd.conditional_formatting.add("B31:F35", FormulaRule(
    formula=['AND(B31<>"",B31>=$H31)'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"])))
kd.conditional_formatting.add("B31:F35", FormulaRule(
    formula=['AND(B31<>"",B31<$G31)'], fill=red,
    font=Font(name=FONT, bold=True, color=PALETTE["bad_text"])))
kd.conditional_formatting.add("B31:F35", FormulaRule(
    formula=['AND(B31<>"",B31>=$G31,B31<$H31)'], fill=amber,
    font=Font(name=FONT, bold=True, color=PALETTE["warn_text"])))
metric_row(38, "Cash per 100 DMs Sent", lambda L: f'IFERROR({L}27/{L}19*100,"")',
           fmt='"$"#,##0', bold=True)
kd["A38"].comment = Comment(
    "The number that makes the daily standards feel worth it: what a hundred DMs is "
    "actually worth to you in cash.", "The Called")

# band stops at F so the Target column keeps its own header cell
section(kd, 40, "DAILY STANDARDS — are you holding the line?", 6)
tg = kd["G40"]
tg.value = "Target"
tg.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
tg.fill = PatternFill("solid", fgColor=PALETTE["dark"])
tg.alignment = Alignment(horizontal="center")
for row, lbl_text, mrow, sref in ((41, "Avg Reels / day", "15", S_REELS),
                                  (42, "Avg Stories / day", "16", S_STORIES),
                                  (43, "Avg Opener DMs / day", "17", S_OPENER),
                                  (44, "Avg Follow-Up DMs / day", "18", S_FOLLOW)):
    metric_row(row, lbl_text, lambda L, m=mrow: f'IFERROR({L}{m}/{L}14,"")', fmt="0.0")
    g = kd[f"G{row}"]
    g.value = f"=Setup!${sref[0]}${sref[1:]}"
    g.number_format = "0.0"
    g.font = Font(name=FONT, size=10, bold=True, color=PALETTE["dark"])
    g.alignment = Alignment(horizontal="center")
    g.fill = PatternFill("solid", fgColor=PALETTE["band_soft"])
    g.border = BOX
metric_row(45, "Perfect Days (all 4 standards hit)",
           lambda L: (f'COUNTIFS({DATES},">="&{L}$10,{DATES},"<="&{L}$11,{col("O")},4)'))
metric_row(46, "Standards Hit Rate", lambda L: f'IFERROR({L}45/{L}14,"")',
           fmt="0.0%", bold=True)

kd.conditional_formatting.add("B41:F44", FormulaRule(
    formula=['AND(B41<>"",B41>=$G41)'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"])))
kd.conditional_formatting.add("B41:F44", FormulaRule(
    formula=['AND(B41<>"",B41<$G41)'], fill=red,
    font=Font(name=FONT, bold=True, color=PALETTE["bad_text"])))

# ---------------------------------------------------------------- chart styling
def _text_props(size=800, color=None, bold=False):
    cp = CharacterProperties(sz=size, b=bold, solidFill=color or PALETTE["muted"])
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def style_chart(chart, title, colors, kind="line"):
    """Flat, gridline-free, brand-coloured. Excel's presets are what 'ancient' looks like."""
    chart.style = None
    chart.title = title
    try:
        chart.title.tx.rich.p[0].r[0].rPr = CharacterProperties(
            sz=1100, b=True, solidFill=PALETTE["ink"])
    except Exception:
        pass
    chart.y_axis.majorGridlines = None
    chart.x_axis.majorGridlines = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.txPr = _text_props()
    chart.y_axis.txPr = _text_props()
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=PALETTE["line"]))
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    no_border = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.spPr = no_border
    if len(chart.series) > 1:
        chart.legend.position = "b"
        chart.legend.overlay = False
        chart.legend.txPr = _text_props(size=900)
    else:
        chart.legend = None
    for s, color in zip(chart.series, colors):
        if kind == "line":
            s.graphicalProperties = GraphicalProperties(
                ln=LineProperties(solidFill=color, w=28575))
            s.smooth = False
            s.marker = Marker(symbol="circle", size=6,
                              spPr=GraphicalProperties(
                                  solidFill=color, ln=LineProperties(solidFill=color)))
        else:
            s.graphicalProperties = GraphicalProperties(
                solidFill=color, ln=LineProperties(noFill=True))
    if kind == "bar":
        chart.gapWidth = 45
    chart.height = 8.4
    chart.width = 15.5
    return chart


def chart_ref(sheet, c1, c2=None, first=None, last=None):
    return Reference(sheet, min_col=c1, max_col=c2 or c1,
                     min_row=first, max_row=last)


section(kd, 48, "TRENDS — your last 13 weeks", 11)
cats = Reference(wr, min_col=2, min_row=WR_FIRST, max_row=WR_LAST)

ch1 = LineChart()
ch1.add_data(chart_ref(wr, 8, 9, 3, WR_LAST), titles_from_data=True)
ch1.set_categories(cats)
kd.add_chart(style_chart(ch1, "DMs sent vs replies", CHART_SERIES[:2], "line"), "A49")

ch2 = BarChart()
ch2.type = "col"
ch2.add_data(chart_ref(wr, 11, 11, 3, WR_LAST), titles_from_data=True)
ch2.add_data(chart_ref(wr, 13, 13, 3, WR_LAST), titles_from_data=True)
ch2.set_categories(cats)
kd.add_chart(style_chart(ch2, "Calls booked vs closes", CHART_SERIES[:2], "bar"), "H49")

ch3 = BarChart()
ch3.type = "col"
ch3.add_data(chart_ref(wr, 14, 14, 3, WR_LAST), titles_from_data=True)
ch3.set_categories(cats)
kd.add_chart(style_chart(ch3, "Cash collected per week", [CHART_SERIES[0]], "bar"), "A71")

ch4 = LineChart()
for c_ in (16, 18, 20):
    ch4.add_data(chart_ref(wr, c_, c_, 3, WR_LAST), titles_from_data=True)
ch4.set_categories(cats)
kd.add_chart(style_chart(ch4, "Reply, book and close rate", CHART_SERIES[:3], "line"), "H71")

# ================================================================ INSTAGRAM
ig = wb.create_sheet("Instagram Tracker")
ig.sheet_view.showGridLines = False
IG_FIRST = 20
IG_LAST = IG_FIRST + IG_WEEKS - 1
banner(ig, "INSTAGRAM TRACKER", 14)
subtitle(ig, "Weekly numbers off your Professional Dashboard, plus a monthly screenshot for your coach to audit.", 14)
widths(ig, {"A": 14, "B": 11, "C": 10, "D": 9, "E": 13, "F": 11, "G": 11, "H": 11,
            "I": 10, "J": 9, "K": 9, "L": 10, "M": 13, "N": 30})

section(ig, 4, "MONTHLY PROFILE CHECK — first Sunday of the month", 14)
for i, s in enumerate([
        "1.  Screenshot your profile — bio, follower count and the top of your grid.",
        "2.  Screenshot Professional Dashboard → Insights (last 30 days).",
        "3.  Click a box below, then Insert → Image → Image in cell, and upload."]):
    c = ig.cell(row=5 + i, column=1, value=s)
    ig.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=8)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(horizontal="left", indent=1)

drop_border = Border(*(Side(style="medium", color=PALETTE["band"]),) * 4)
for start_c, end_c, title in ((1, 5, "PROFILE SCREENSHOT"), (7, 11, "INSIGHTS SCREENSHOT")):
    t = ig.cell(row=9, column=start_c, value=title)
    t.font = Font(name=FONT, size=9, bold=True, color=PALETTE["ink"])
    ig.merge_cells(start_row=10, start_column=start_c, end_row=16, end_column=end_c)
    box = ig.cell(row=10, column=start_c, value="drop image here")
    box.font = Font(name=FONT, size=10, italic=True, color=PALETTE["muted"])
    box.fill = PatternFill("solid", fgColor=PALETTE["band_soft"])
    box.alignment = Alignment(horizontal="center", vertical="center")
    for rr in range(10, 17):
        for cc in range(start_c, end_c + 1):
            ig.cell(row=rr, column=cc).border = drop_border

section(ig, 18, "WEEKLY NUMBERS — log every Sunday night", 14)
header_row(ig, 19, ["Week Starting", "Followers\n(end of week)", "Growth\n(+/-)", "Growth\n%",
                    "Accounts\nReached", "Interactions", "Profile\nVisits", "New\nFollowers",
                    "Link\nClicks", "Posts\n(auto)", "Stories\n(auto)", "Reach\nper Post",
                    "New Followers\nper 1k Reached", "Notes"])
ig.freeze_panes = "B20"

ig_inputs = list("BEFGHI") + ["N"]
for i in range(IG_WEEKS):
    r = IG_FIRST + i
    a = ig.cell(row=r, column=1,
                value=f"=Setup!${S_WEEK1[0]}${S_WEEK1[1:]}" if i == 0 else f"=A{r-1}+7")
    a.number_format = "ddd d mmm"
    a.font = Font(name=FONT, size=10)
    a.fill = auto_fill
    a.border = BOX
    for letter in ig_inputs:
        c = ig[f"{letter}{r}"]
        c.fill = input_fill
        c.font = Font(name=FONT, size=10)
        c.border = BOX
        c.alignment = Alignment(horizontal="center" if letter != "N" else "left")
        if letter != "N":
            c.number_format = "#,##0"
    if i > 0:
        ig.cell(row=r, column=3, value=f'=IF(OR(B{r}="",B{r-1}=""),"",B{r}-B{r-1})')
        ig.cell(row=r, column=4,
                value=f'=IF(OR(B{r}="",B{r-1}=""),"",IFERROR((B{r}-B{r-1})/B{r-1},""))')
    for cidx, key in ((10, "reels"), (11, "stories")):
        ig.cell(row=r, column=cidx, value=(
            f'=IF($A{r}>TODAY(),"",SUMIFS({col(M[key])},{DATES},">="&$A{r},'
            f'{DATES},"<="&$A{r}+6))'))
    ig.cell(row=r, column=12, value=f'=IFERROR(E{r}/J{r},"")')
    ig.cell(row=r, column=13, value=f'=IFERROR(H{r}/E{r}*1000,"")')
    for cidx, fmt in ((3, "+#,##0;-#,##0;0"), (4, "0.0%"), (10, "#,##0"), (11, "#,##0"),
                      (12, "0.0"), (13, "0.0")):
        c = ig.cell(row=r, column=cidx)
        c.number_format = fmt
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="center")
        c.fill = auto_fill
        c.border = BOX

dv_ig = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                       allow_blank=True, showErrorMessage=True,
                       errorTitle="Whole numbers only",
                       error="Enter a whole number of 0 or more.")
ig.add_data_validation(dv_ig)
for letter in "BEFGHI":
    dv_ig.add(f"{letter}{IG_FIRST}:{letter}{IG_LAST}")

ig.conditional_formatting.add(f"C{IG_FIRST}:D{IG_LAST}", FormulaRule(
    formula=[f'AND($C{IG_FIRST}<>"",$C{IG_FIRST}>0)'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"])))
ig.conditional_formatting.add(f"C{IG_FIRST}:D{IG_LAST}", FormulaRule(
    formula=[f'AND($C{IG_FIRST}<>"",$C{IG_FIRST}<0)'], fill=red,
    font=Font(name=FONT, bold=True, color=PALETTE["bad_text"])))

ig_cats = Reference(ig, min_col=1, min_row=IG_FIRST, max_row=IG_LAST)
ic1 = LineChart()
ic1.add_data(chart_ref(ig, 2, 2, 19, IG_LAST), titles_from_data=True)
ic1.set_categories(ig_cats)
ig.add_chart(style_chart(ic1, "Follower growth", [CHART_SERIES[0]], "line"), "P4")

ic2 = BarChart()
ic2.type = "col"
ic2.add_data(chart_ref(ig, 5, 5, 19, IG_LAST), titles_from_data=True)
ic2.set_categories(ig_cats)
ig.add_chart(style_chart(ic2, "Accounts reached per week", [CHART_SERIES[1]], "bar"), "P22")

ig["J19"].comment = Comment(
    "Posts and Stories fill in automatically from your Daily Log — don't type here.",
    "The Called")

# ================================================================ CHEAT SHEET
cs = wb.create_sheet("Cheat Sheet")
cs.sheet_view.showGridLines = False
widths(cs, {"A": 24, "B": 104, "C": 14, "D": 16})
banner(cs, "CHEAT SHEET — what your numbers are telling you", 4)
subtitle(cs, "Every number in this sheet has one job: to tell you which single thing to fix next.", 4)

section(cs, 4, "HOW TO READ YOUR NUMBERS", 4)
for i, t in enumerate(INTRO):
    c = cs.cell(row=5 + i, column=2, value=t)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    cs.row_dimensions[5 + i].height = 28

section(cs, 9, "HEALTHY RANGES — edit these as you learn what's true for your clients", 4)
header_row(cs, 10, ["Rate", "What it's measuring", "Red\nbelow", "The bar\n(green at/above)"])
for i, rate in enumerate(RATES):
    r = 11 + i
    label(cs, r, rate["name"], bold=True)
    cs[f"A{r}"].border = BOX
    w = cs.cell(row=r, column=2, value=f'{rate["formula"]} — {rate["measures"]}')
    w.font = Font(name=FONT, size=10)
    w.alignment = Alignment(vertical="center", indent=1)
    w.border = BOX
    for cidx, v in ((3, rate["needs_work"]), (4, rate["bar"])):
        c = cs.cell(row=r, column=cidx, value=v)
        c.number_format = "0%"
        c.font = Font(name=FONT, size=11, bold=True)
        c.fill = input_fill
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
    cs.row_dimensions[r].height = 20

n = cs.cell(row=16, column=2, value=(
    "Two numbers, two meanings. THE BAR is what we push you towards — hit it and the dashboard "
    "goes green. RED BELOW is the line where something is actually broken. In between you are "
    "amber: working, not there yet. That is most people most of the time, and it is not a "
    "failure. Both are cream cells because your coach can tune them to you, and the KPI "
    "Dashboard re-grades itself the moment they change."))
n.font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])
n.alignment = Alignment(wrap_text=True, vertical="top")
cs.row_dimensions[16].height = 32


def block(row, title, items):
    """One metric: a banded title, then label/explanation pairs."""
    cs.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    t = cs.cell(row=row, column=1, value=title)
    t.font = Font(name=FONT, size=11, bold=True, color=PALETTE["dark"])
    t.fill = PatternFill("solid", fgColor=PALETTE["band_soft"])
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cs.row_dimensions[row].height = 20
    r = row + 1
    for lbl_text, text in items:
        if not text or text == "—":
            continue
        lc = cs.cell(row=r, column=1, value=lbl_text)
        lc.font = Font(name=FONT, size=9, bold=True,
                       color=PALETTE["bad_text"] if "low" in lbl_text.lower()
                       else PALETTE["good_text"] if "high" in lbl_text.lower()
                       else PALETTE["ink"])
        lc.alignment = Alignment(horizontal="right", vertical="top", indent=1)
        tc = cs.cell(row=r, column=2, value=text)
        tc.font = Font(name=FONT, size=10)
        tc.alignment = Alignment(wrap_text=True, vertical="top")
        cs.row_dimensions[r].height = max(16, (len(text) // 100 + 1) * 13 + 4)
        r += 1
    return r + 1


section(cs, 18, "THE FIVE RATES — where the money actually leaks", 4)
r = 19
for rate in RATES:
    r = block(r, f'{rate["name"].upper()} — {rate["formula"].lower()}', [
        ("Measures", rate["measures"]),
        ("If it's low", rate["low"]),
        ("Do this", rate["do"]),
        ("If it's high", rate["high"]),
    ])

section(cs, r, "THE VOLUME NUMBERS — the inputs you control", 4)
r += 1
for v in VOLUME:
    r = block(r, v["name"].upper(), [
        ("What it does", v["does"]), ("If it's low", v["low"]), ("Do this", v["do"])])

section(cs, r, "THE MONEY NUMBERS", 4)
r += 1
for m in MONEY:
    r = block(r, m["name"].upper(), [("What it is", m["what"]), ("Use it", m["use"])])

section(cs, r, "THE INSTAGRAM NUMBERS", 4)
r += 1
for g in INSTAGRAM:
    r = block(r, g["name"].upper(), [("What it is", g["what"]), ("Watch", g["watch"])])

section(cs, r, "FIVE TRAPS", 4)
r += 1
for i, t in enumerate(TRAPS):
    c = cs.cell(row=r + i, column=1, value=f"{i + 1}.")
    c.font = Font(name=FONT, size=11, bold=True, color=PALETTE["ink"])
    c.alignment = Alignment(horizontal="right", vertical="top", indent=1)
    tc = cs.cell(row=r + i, column=2, value=t)
    tc.font = Font(name=FONT, size=10)
    tc.alignment = Alignment(wrap_text=True, vertical="top")
    cs.row_dimensions[r + i].height = max(18, (len(t) // 100 + 1) * 13 + 5)

# ================================================================ finish
order = ["Start Here", "Setup", "Daily Log", "This Week", "KPI Dashboard",
         "Instagram Tracker", "Cheat Sheet", "Weekly Rollup"]
wb._sheets = [wb[n] for n in order]
for n in order:
    wb[n].sheet_properties.tabColor = (
        PALETTE["tab_input"] if n in ("Setup", "Daily Log", "Instagram Tracker")
        else PALETTE["ink"] if n == "Cheat Sheet"
        else PALETTE["band"] if n == "This Week"
        else PALETTE["dark"])
wb.active = 0

DEMO = "--demo" in sys.argv
DEMO_WEEKS = 17
DEMO_DAYS = DEMO_WEEKS * 7


def demo_phase(week):
    """A story, not noise: a solid start, follow-ups collapse and take booked calls
    with them, then a fix and a recovery. Every chart gets shape, and the
    'fix this first' callout fires on a cause you can actually point at."""
    if week <= 4:                      # finding their feet, holding the standard
        return dict(reels=(5, 7), stories=(5, 7), opener=(20, 26), follow=(10, 13),
                    reply=0.19, pitch=0.55, book=0.31, show=0.74, close=0.27)
    if week <= 8:                      # follow-ups collapse — everything downstream follows
        return dict(reels=(4, 6), stories=(3, 5), opener=(18, 24), follow=(1, 4),
                    reply=0.11, pitch=0.44, book=0.21, show=0.63, close=0.21)
    if week <= 10:                     # caught it, starting to fix it
        return dict(reels=(5, 7), stories=(4, 6), opener=(20, 25), follow=(6, 9),
                    reply=0.15, pitch=0.50, book=0.26, show=0.70, close=0.24)
    if week < DEMO_WEEKS - 1:
        return dict(reels=(6, 8), stories=(5, 7), opener=(23, 29), follow=(11, 15),
                    reply=0.21, pitch=0.59, book=0.35, show=0.78, close=0.30)
    # final week: the work is up but the book rate quietly fell. This is the
    # story worth demoing — the standards all read green and the sheet still
    # finds the thing costing them money.
    return dict(reels=(6, 8), stories=(5, 7), opener=(23, 29), follow=(11, 15),
                reply=0.21, pitch=0.57, book=0.16, show=0.76, close=0.28)


if DEMO:
    rnd = random.Random(11)
    st["C5"] = "Demo Client"
    st["C6"] = "Francis"
    st["C7"] = "@democlient"
    st["C8"] = f"=TODAY()-{DEMO_DAYS - 1}"   # last row lands on TODAY
    first_day = datetime.date.today() - datetime.timedelta(days=DEMO_DAYS - 1)

    followers = 3180
    prev_booked = 0
    blank_days = {38, 73}

    def jitter(x, spread=0.35):
        return max(0, int(round(x * (1 + rnd.uniform(-spread, spread)))))

    for i in range(DEMO_DAYS):
        r = DL_FIRST + i
        ph = demo_phase(i // 7)
        weekend = (first_day + datetime.timedelta(days=i)).weekday() >= 5
        if i in blank_days:
            vals = {k: 0 for k in ("reels", "stories", "opener", "follow", "replies",
                                   "pitched", "booked", "showed", "closes", "cash", "revenue")}
        else:
            scale = 0.35 if weekend else 1.0
            vals = {}
            for key in ("reels", "stories", "opener", "follow"):
                lo, hi = ph[key]
                vals[key] = max(0, int(round(rnd.randint(lo, hi) * scale)))
            dms = vals["opener"] + vals["follow"]
            vals["replies"] = jitter(dms * ph["reply"])
            vals["pitched"] = jitter(vals["replies"] * ph["pitch"])
            vals["booked"] = jitter(vals["pitched"] * ph["book"])
            vals["showed"] = 0
            vals["closes"] = 0
            vals["cash"] = 0
            vals["revenue"] = 0
        # calls booked yesterday show up (or don't) today — a real show rate
        vals["showed"] = min(prev_booked,
                             sum(1 for _ in range(prev_booked) if rnd.random() < ph["show"]))
        prev_booked = vals["booked"]
        for _ in range(vals["showed"]):
            if rnd.random() < ph["close"]:
                vals["closes"] += 1
                amount = rnd.choice([1500, 2000, 2000, 3000])
                vals["cash"] += amount
                vals["revenue"] += amount + rnd.choice([0, 0, 1500])
        if i >= DEMO_DAYS - 3 and i not in blank_days:
            # finish on a run of perfect days so the streak tile isn't a zero on camera
            vals["reels"] = max(vals["reels"], 6)
            vals["stories"] = max(vals["stories"], 6)
            vals["opener"] = max(vals["opener"], 22)
            vals["follow"] = max(vals["follow"], 11)
        followers += 0 if i in blank_days else rnd.randint(
            4 if 5 <= i // 7 <= 8 else 9, 14 if 5 <= i // 7 <= 8 else 32)
        for key, letter in M.items():
            if key == "followers":
                continue
            dl[f"{letter}{r}"] = vals[key]
        dl[f"N{r}"] = followers
    dl["P4"] = "First week — finding the rhythm"
    dl[f"P{DL_FIRST + 35}"] = "Slipping on follow-ups, got busy with delivery"
    dl[f"P{DL_FIRST + 38}"] = "Travel day, nothing done"
    dl[f"P{DL_FIRST + 70}"] = "Coach flagged follow-ups on our call — fixing it"
    dl[f"P{DL_FIRST + 84}"] = "Back on follow-ups properly this week"

    # Instagram: reach sags through the slump, then climbs past where it started
    reach_arc = [9800, 11200, 12600, 13900, 15200, 11000, 9400, 8600, 9100,
                 11800, 13200, 15600, 17400, 19100, 21000, 23500, 26800]
    fol = 3180
    for w, reach in enumerate(reach_arc):
        r = IG_FIRST + w
        gain = int(reach * rnd.uniform(0.006, 0.009))
        ig[f"B{r}"] = fol + gain
        ig[f"E{r}"] = reach
        ig[f"F{r}"] = int(reach * rnd.uniform(0.045, 0.055))
        ig[f"G{r}"] = int(reach * rnd.uniform(0.028, 0.034))
        ig[f"H{r}"] = gain
        ig[f"I{r}"] = int(reach * rnd.uniform(0.002, 0.003))
        fol += gain
    ig[f"N{IG_FIRST + 6}"] = "Posted less this month — reach went with it"
    ig[f"N{IG_FIRST + 11}"] = "Back to 5 a day, reach recovering"

    # Unmissable on every tab: this copy is a reference, never a working sheet.
    # Marking row 1 rather than inserting a warning row keeps every formula intact.
    ws["B4"] = "This is a sample. Do not work in it."
    ws["B5"] = ("Every number in this file is made up, and it is here so you can see what the "
                "Scoreboard looks like once it has history in it. Your own copy starts empty — "
                "log in that one, not this one.")
    for sheet in wb.worksheets:
        head = sheet["A1"]
        head.value = f"DEMO \u00b7 {head.value} \u00b7 SAMPLE DATA — REFERENCE ONLY"
        head.fill = PatternFill("solid", fgColor=PALETTE["clay"])
        sheet.sheet_properties.tabColor = PALETTE["clay"]

out = ("docs/The_Called_Scoreboard_DEMO.xlsx" if DEMO
       else "docs/The_Called_Scoreboard.xlsx")
wb.save(out)


def notion_cheat_sheet(path):
    """Same content as the Cheat Sheet tab, as markdown tables that paste
    straight into Notion. Generated from the same module so the two can't drift."""
    def esc(t):
        return t.replace("|", "\\|")

    L = ["# Reading Your Scoreboard", ""]
    L += [esc(p_) + "\n" for p_ in INTRO]
    L += ["## The five rates at a glance", "",
          "| Rate | Formula | The bar | Red below | What it measures |",
          "| --- | --- | --- | --- | --- |"]
    for x in RATES:
        L.append(f'| **{x["name"]}** | {x["formula"]} | {x["bar"]:.0%} | {x["needs_work"]:.0%} '
                 f'| {esc(x["measures"])} |')
    L += ["", "*Two numbers, two meanings. **The bar** is what we push you towards — hit it and "
              "your dashboard goes green. **Red below** is where something is actually broken. "
              "In between you are amber: working, not there yet. That is most people most of the "
              "time, and it is not a failure. Both live in editable cells on the Cheat Sheet tab "
              "of your Scoreboard, and your coach can tune them to you.*", "",
          "## Diagnosing a rate", "",
          "| Rate | If it's low, it usually means | Do this | If it's high |",
          "| --- | --- | --- | --- |"]
    for x in RATES:
        L.append(f'| **{x["name"]}** | {esc(x["low"])} | {esc(x["do"])} | {esc(x["high"])} |')
    L += ["", "## The volume numbers — the inputs you control", "",
          "| Number | What it does | If it's low | Do this |", "| --- | --- | --- | --- |"]
    for x in VOLUME:
        do = "" if x["do"] == "—" else esc(x["do"])
        L.append(f'| **{x["name"]}** | {esc(x["does"])} | {esc(x["low"])} | {do} |')
    L += ["", "## The money numbers", "",
          "| Number | What it is | How to use it |", "| --- | --- | --- |"]
    for x in MONEY:
        L.append(f'| **{x["name"]}** | {esc(x["what"])} | {esc(x["use"])} |')
    L += ["", "## The Instagram numbers", "",
          "| Number | What it is | What to watch |", "| --- | --- | --- |"]
    for x in INSTAGRAM:
        L.append(f'| **{x["name"]}** | {esc(x["what"])} | {esc(x["watch"])} |')
    L += ["", "## Five traps", ""]
    L += [f"{i + 1}. {esc(t)}" for i, t in enumerate(TRAPS)]
    L.append("")
    with open(path, "w") as fh:
        fh.write("\n".join(L))
    return path


if not DEMO:
    print(f"wrote {notion_cheat_sheet('docs/Cheat_Sheet_for_Notion.md')}")
print(f"wrote {out}")
print(f"  logo: {LOGO_PATH or 'none yet — drop a PNG at ' + LOGO_DEFAULT + ' and re-run'}")
print(f"  Daily Log rows {DL_FIRST}-{DL_LAST} ({DAYS} days)")
print(f"  Weekly Rollup rows {WR_FIRST}-{WR_LAST} (rolling {ROLL_WEEKS} weeks)")
print(f"  Instagram rows {IG_FIRST}-{IG_LAST} ({IG_WEEKS} weeks)")
