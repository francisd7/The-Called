#!/usr/bin/env python3
"""Build The Called — Client Tracker workbook (.xlsx, imports cleanly into Google Sheets).

Run:  python3 scripts/build_client_tracker.py
Out:  docs/The_Called_Client_Tracker.xlsx

Brand colours live in PALETTE below — change the hex codes and re-run to recolour
the whole workbook.

Design rules that matter (don't undo them without reading this):
  * No Sheets-only functions (QUERY/ARRAYFORMULA/SPARKLINE) — file must open in both.
  * Conditional formatting NEVER references another tab: Google Sheets forbids it.
    Targets are mirrored onto the tab that needs them (Daily Log row 2, Dashboard col G).
  * Weekly Rollup is a rolling LAST 13 WEEKS, not a fixed calendar — keeps dashboard
    charts always-current with no empty tail.
"""

import sys
import random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

# ---------------------------------------------------------------- brand
PALETTE = {
    "dark":        "111318",  # headers / banners
    "accent":      "6B4EFF",  # primary brand accent
    "accent_soft": "EFEAFF",  # accent tint for section strips
    "input":       "FFF8DC",  # "type here" cells
    "auto":        "F1F3F5",  # "leave alone" cells
    "good":        "E6F4EA",
    "good_text":   "137333",
    "bad":         "FCE8E6",
    "bad_text":    "C5221F",
    "line":        "D7DAE0",
    "muted":       "6B7280",
}
FONT = "Arial"

DAYS = 366          # rows in the Daily Log
IG_WEEKS = 53       # rows in the Instagram weekly log
ROLL_WEEKS = 13     # rolling window on the Weekly Rollup

DL_FIRST, DL_LAST = 4, 4 + DAYS - 1           # Daily Log data rows
DATES = f"'Daily Log'!$A${DL_FIRST}:$A${DL_LAST}"
LOGGED = f"'Daily Log'!$Q${DL_FIRST}:$Q${DL_LAST}"


def col(letter):
    return f"'Daily Log'!${letter}${DL_FIRST}:${letter}${DL_LAST}"


# Daily Log metric columns
M = {"reels": "C", "stories": "D", "opener": "E", "follow": "F", "replies": "G",
     "pitched": "H", "booked": "I", "showed": "J", "closes": "K",
     "cash": "L", "revenue": "M", "followers": "N"}

# ---------------------------------------------------------------- helpers
thin = Side(style="thin", color=PALETTE["line"])
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def banner(ws, text, last_col, row=1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=PALETTE["dark"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 30


def subtitle(ws, text, last_col, row=2):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10, italic=True, color=PALETTE["muted"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 18


def section(ws, row, text, last_col):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10, bold=True, color=PALETTE["accent"])
    c.fill = PatternFill("solid", fgColor=PALETTE["accent_soft"])
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)


def header_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PALETTE["dark"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 32


def label(ws, row, text, bold=False, indent=1, colnum=1):
    c = ws.cell(row=row, column=colnum, value=text)
    c.font = Font(name=FONT, size=10, bold=bold)
    c.alignment = Alignment(horizontal="left", indent=indent)
    return c


def widths(ws, mapping):
    for k, v in mapping.items():
        ws.column_dimensions[k].width = v


def sumifs(metric_col, start_ref, end_ref):
    return (f"SUMIFS({col(metric_col)},{DATES},\">=\"&{start_ref},"
            f"{DATES},\"<=\"&{end_ref})")


wb = Workbook()

# ================================================================ START HERE
ws = wb.active
ws.title = "Start Here"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 3, "B": 26, "C": 58, "D": 26, "E": 14, "F": 14, "G": 14, "H": 14})

banner(ws, "THE CALLED — CLIENT TRACKER", 8)
subtitle(ws, "One sheet. Fill it in daily, and your numbers tell you exactly what to fix.", 8)

ws["B4"] = "What this is"
ws["B4"].font = Font(name=FONT, size=12, bold=True)
ws["B5"] = ("Your whole game in one place: the work you put in (posts, stories, DMs) and what it "
            "turns into (replies, calls, closes, cash). Everything on the KPI Dashboard is "
            "calculated for you — you never type a number twice.")
ws.merge_cells("B5:H7")
ws["B5"].alignment = Alignment(wrap_text=True, vertical="top")
ws["B5"].font = Font(name=FONT, size=10)

section(ws, 9, "YOUR THREE JOBS", 8)
jobs = [
    ("Every day (60 seconds)", "Open Daily Log, find today's row, fill in what you did."),
    ("Every Sunday (2 minutes)", "Open Instagram Tracker, log this week's numbers from your Professional Dashboard."),
    ("Every month (1 minute)", "Drop a profile + insights screenshot at the top of Instagram Tracker so your coach can audit your page."),
]
r = 10
for when, what in jobs:
    label(ws, r, when, bold=True)
    ws.cell(row=r, column=3, value=what).font = Font(name=FONT, size=10)
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 28
    r += 1

section(ws, 14, "WHAT EACH TAB IS FOR", 8)
header_row(ws, 15, ["Tab", "What it's for", "Do you type here?"], start_col=2)
tabs = [
    ("Daily Log", "One row per day. The only place your daily numbers get entered.", "YES — every day"),
    ("KPI Dashboard", "Auto totals, conversion rates and charts for this week, this month and all time.", "No — it's automatic"),
    ("Instagram Tracker", "Weekly follower and reach numbers, plus your monthly profile screenshots.", "YES — weekly"),
    ("Weekly Rollup", "Feeds the dashboard charts — your last 13 weeks, week by week.", "No — it's automatic"),
    ("Setup", "Your name, start date and daily standards. Set once at the start.", "YES — once"),
]
r = 16
for t, w, y in tabs:
    ws.cell(row=r, column=2, value=t).font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=3, value=w).font = Font(name=FONT, size=10)
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
    c = ws.cell(row=r, column=4, value=y)
    c.font = Font(name=FONT, size=10, bold=y.startswith("YES"),
                  color=PALETTE["accent"] if y.startswith("YES") else PALETTE["muted"])
    for cc in range(2, 5):
        ws.cell(row=r, column=cc).border = BOX
    ws.row_dimensions[r].height = 26
    r += 1

section(ws, 22, "THE RULES (read once, saves you a headache)", 8)
rules = [
    "Cream cells = you type here. Grey cells = calculated, leave them alone.",
    "Don't delete columns or rows — it breaks the maths. If you don't use a column, right-click and Hide it.",
    "One row per day, in date order. Log a zero day as 0, not blank — blank means 'didn't track'.",
    "Miss a day? Backfill it. A gap makes your averages look better than they are.",
]
r = 23
for x in rules:
    ws.cell(row=r, column=2, value="•").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=3, value=x).font = Font(name=FONT, size=10)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
    r += 1

section(ws, 28, "EXAMPLE — what one good day looks like in the Daily Log", 8)
ex_head = ["Date", "Reels", "Stories", "Opener DMs", "Follow-Ups", "Replies", "Pitched"]
ex_vals = ["Mon 15 Sep", 4, 30, 10, 10, 6, 3]
for i, (h, v) in enumerate(zip(ex_head, ex_vals)):
    hc = ws.cell(row=29, column=2 + i, value=h)
    hc.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
    hc.fill = PatternFill("solid", fgColor=PALETTE["dark"])
    hc.alignment = Alignment(horizontal="center")
    vc = ws.cell(row=30, column=2 + i, value=v)
    vc.font = Font(name=FONT, size=10)
    vc.fill = PatternFill("solid", fgColor=PALETTE["good"])
    vc.alignment = Alignment(horizontal="center")
    vc.border = BOX
ws["B31"] = "All four standards hit → that day scores 4/4 and turns green."
ws["B31"].font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])

section(ws, 33, "WALKTHROUGH VIDEOS", 8)
vids = [("Daily Log walkthrough", "[paste Loom link]"),
        ("KPI Dashboard walkthrough", "[paste Loom link]"),
        ("Instagram Tracker walkthrough", "[paste Loom link]")]
r = 34
for n, link in vids:
    label(ws, r, n, bold=True)
    c = ws.cell(row=r, column=3, value=link)
    c.font = Font(name=FONT, size=10, color=PALETTE["accent"])
    c.fill = PatternFill("solid", fgColor=PALETTE["input"])
    c.border = BOX
    r += 1

# ================================================================ SETUP
st = wb.create_sheet("Setup")
st.sheet_view.showGridLines = False
widths(st, {"A": 3, "B": 34, "C": 20, "D": 52})
banner(st, "SETUP — fill this in once", 4)
subtitle(st, "Cream cells are yours to fill. Everything in the workbook reads from here.", 4)


def setup_row(row, lbl, value, note, fmt=None, is_input=True):
    label(st, row, lbl, bold=True, colnum=2)
    c = st.cell(row=row, column=3, value=value)
    c.font = Font(name=FONT, size=10)
    c.fill = PatternFill("solid", fgColor=PALETTE["input"] if is_input else PALETTE["auto"])
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
    if fmt:
        c.number_format = fmt
    n = st.cell(row=row, column=4, value=note)
    n.font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])
    return c


section(st, 4, "CLIENT DETAILS", 4)
setup_row(5, "Client name", "", "Your name — shows on the dashboard.")
setup_row(6, "Coach / CSM", "", "Who you report to inside The Called.")
setup_row(7, "Instagram handle", "", "e.g. @yourhandle")
setup_row(8, "Tracking start date", "=TODAY()", "Day 1. This sets every date in the Daily Log.",
          fmt="ddd d mmm yyyy")
setup_row(9, "First tracking week (auto)", "=C8-WEEKDAY(C8,2)+1",
          "Monday of your start week. Calculated — don't edit.",
          fmt="ddd d mmm yyyy", is_input=False)

section(st, 11, "DAILY STANDARDS — the minimum you hold yourself to", 4)
setup_row(12, "Reels / posts per day", 4, "Default 4.")
setup_row(13, "Stories per day", 30, "Default 30.")
setup_row(14, "Opener DMs per day", 10, "Default 10.")
setup_row(15, "Follow-up DMs per day", 10, "Default 10.")
st["D12"].comment = Comment(
    "Defaults are The Called's standard: 4 posts, 30 stories, 10 opener DMs, 10 follow-up DMs. "
    "Change them per client if their programme says something different.", "The Called")

section(st, 17, "GOALS", 4)
setup_row(18, "Monthly cash collected goal", 10000, "Drives the goal box on the dashboard.", fmt='"$"#,##0')
setup_row(19, "Your offer price", 0, "Optional — for your own maths.", fmt='"$"#,##0')

# ================================================================ DAILY LOG
dl = wb.create_sheet("Daily Log")
dl.sheet_view.showGridLines = False
banner(dl, "DAILY LOG — one row per day", 17)

# row 2 mirrors the targets locally (Google Sheets conditional formatting can't
# reference another tab — this is what the green/red rules compare against)
lab = dl.cell(row=2, column=1, value="Daily standard →")
lab.font = Font(name=FONT, size=9, bold=True, color=PALETTE["accent"])
lab.alignment = Alignment(horizontal="right")
dl.merge_cells("H2:P2")
note = dl["H2"]
note.value = "Fill in what you actually did. Log a zero as 0 — blank means you didn't track that day."
note.font = Font(name=FONT, size=9, italic=True, color=PALETTE["muted"])
note.alignment = Alignment(horizontal="left", vertical="center", indent=1)
dl.row_dimensions[2].height = 18
for c_, ref in (("C", "C12"), ("D", "C13"), ("E", "C14"), ("F", "C15")):
    cell = dl[f"{c_}2"]
    cell.value = f"=Setup!${ref[0]}${ref[1:]}"
    cell.font = Font(name=FONT, size=9, bold=True, color=PALETTE["accent"])
    cell.fill = PatternFill("solid", fgColor=PALETTE["accent_soft"])
    cell.alignment = Alignment(horizontal="center")

dl_headers = ["Date", "Day", "Reels /\nPosts", "Stories", "Opener\nDMs Sent",
              "Follow-Up\nDMs Sent", "Replies", "Calls\nPitched", "Calls\nBooked",
              "Calls\nShowed", "Closes", "Cash\nCollected", "Revenue\nGenerated",
              "Followers\n(end of day)", "Standards\nHit (of 4)", "Notes",
              "Logged?\n(auto)"]
header_row(dl, 3, dl_headers)
widths(dl, {"A": 13, "B": 6, "C": 8, "D": 8, "E": 9, "F": 10, "G": 9, "H": 9, "I": 9,
            "J": 9, "K": 8, "L": 11, "M": 11, "N": 11, "O": 10, "P": 34, "Q": 9})
dl.freeze_panes = "C4"

input_fill = PatternFill("solid", fgColor=PALETTE["input"])
auto_fill = PatternFill("solid", fgColor=PALETTE["auto"])
money_cols = {"L", "M"}
input_cols = list("CDEFGHIJKLMN")

for i in range(DAYS):
    r = DL_FIRST + i
    a = dl.cell(row=r, column=1, value="=Setup!$C$8" if i == 0 else f"=A{r-1}+1")
    a.number_format = "ddd d mmm"
    a.font = Font(name=FONT, size=10)
    a.fill = auto_fill
    b = dl.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"ddd"))')
    b.font = Font(name=FONT, size=9, color=PALETTE["muted"])
    b.fill = auto_fill
    b.alignment = Alignment(horizontal="center")
    for letter in input_cols:
        c = dl[f"{letter}{r}"]
        c.fill = input_fill
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        c.number_format = '"$"#,##0' if letter in money_cols else "#,##0"
    o = dl.cell(row=r, column=15,
                value=(f'=IF($Q{r}=0,"",IF(C{r}>=C$2,1,0)+IF(D{r}>=D$2,1,0)'
                       f'+IF(E{r}>=E$2,1,0)+IF(F{r}>=F$2,1,0))'))
    o.alignment = Alignment(horizontal="center")
    o.font = Font(name=FONT, size=10, bold=True)
    o.border = BOX
    p = dl.cell(row=r, column=16)
    p.fill = input_fill
    p.font = Font(name=FONT, size=9)
    p.border = BOX
    q = dl.cell(row=r, column=17, value=f'=IF(COUNT(C{r}:M{r})>0,1,0)')
    q.fill = auto_fill
    q.font = Font(name=FONT, size=8, color=PALETTE["muted"])
    q.alignment = Alignment(horizontal="center")

dl["Q3"].comment = Comment(
    "Auto column. 1 = you logged something that day, 0 = you didn't. "
    "The dashboard uses it to work out your averages — don't type in it or delete it.",
    "The Called")

dv_int = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Whole numbers only",
                        error="Enter a whole number of 0 or more.")
dv_money = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                          allow_blank=True, showErrorMessage=True,
                          errorTitle="Amount only",
                          error="Enter a dollar amount of 0 or more.")
dl.add_data_validation(dv_int)
dl.add_data_validation(dv_money)
for letter in "CDEFGHIJKN":
    dv_int.add(f"{letter}{DL_FIRST}:{letter}{DL_LAST}")
for letter in "LM":
    dv_money.add(f"{letter}{DL_FIRST}:{letter}{DL_LAST}")

green = PatternFill("solid", bgColor=PALETTE["good"])
red = PatternFill("solid", bgColor=PALETTE["bad"])
rng = f"C{DL_FIRST}:F{DL_LAST}"
dl.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($Q{DL_FIRST}=1,C{DL_FIRST}>=C$2)'], fill=green, stopIfTrue=False))
dl.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($Q{DL_FIRST}=1,C{DL_FIRST}<C$2)'], fill=red, stopIfTrue=False))
dl.conditional_formatting.add(f"O{DL_FIRST}:O{DL_LAST}", FormulaRule(
    formula=[f'$O{DL_FIRST}=4'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"]), stopIfTrue=False))
dl.conditional_formatting.add(f"O{DL_FIRST}:O{DL_LAST}", FormulaRule(
    formula=[f'AND($Q{DL_FIRST}=1,$O{DL_FIRST}<2)'], fill=red,
    font=Font(name=FONT, bold=True, color=PALETTE["bad_text"]), stopIfTrue=False))

# ================================================================ WEEKLY ROLLUP
wr = wb.create_sheet("Weekly Rollup")
wr.sheet_view.showGridLines = False
WR_FIRST = 4
WR_LAST = WR_FIRST + ROLL_WEEKS - 1
banner(wr, "WEEKLY ROLLUP — your last 13 weeks", 18)
subtitle(wr, "All automatic. This is what the dashboard charts read from. Nothing to fill in here.", 18)

wr_headers = ["Week Starting", "Week", "Days\nLogged", "Reels /\nPosts", "Stories",
              "Opener\nDMs", "Follow-Up\nDMs", "Total DMs\nSent", "Replies",
              "Calls\nPitched", "Calls\nBooked", "Calls\nShowed", "Closes",
              "Cash\nCollected", "Revenue", "Reply\nRate", "Book\nRate", "Close\nRate"]
header_row(wr, 3, wr_headers)
widths(wr, {"A": 14, "B": 9, "C": 8, "D": 8, "E": 8, "F": 9, "G": 9, "H": 9, "I": 9,
            "J": 9, "K": 9, "L": 9, "M": 8, "N": 11, "O": 11, "P": 8, "Q": 8, "R": 8})
wr.freeze_panes = "C4"

GUARD = 'IF($A{r}<Setup!$C$9,"",{body})'
for i in range(ROLL_WEEKS):
    r = WR_FIRST + i
    first = f"=TODAY()-WEEKDAY(TODAY(),2)+1-{7 * (ROLL_WEEKS - 1)}"
    a = wr.cell(row=r, column=1, value=first if i == 0 else f"=A{r-1}+7")
    a.number_format = "ddd d mmm"
    a.font = Font(name=FONT, size=10)
    b = wr.cell(row=r, column=2, value=f'=IF(A{r}="","",TEXT(A{r},"d mmm"))')
    b.font = Font(name=FONT, size=9, color=PALETTE["muted"])
    b.alignment = Alignment(horizontal="center")
    body_days = f'SUMIFS({LOGGED},{DATES},">="&$A{r},{DATES},"<="&$A{r}+6)'
    wr.cell(row=r, column=3, value="=" + GUARD.format(r=r, body=body_days))
    plan = [(4, "reels"), (5, "stories"), (6, "opener"), (7, "follow"), (9, "replies"),
            (10, "pitched"), (11, "booked"), (12, "showed"), (13, "closes"),
            (14, "cash"), (15, "revenue")]
    for cidx, key in plan:
        body = (f'SUMIFS({col(M[key])},{DATES},">="&$A{r},{DATES},"<="&$A{r}+6)')
        wr.cell(row=r, column=cidx, value="=" + GUARD.format(r=r, body=body))
    wr.cell(row=r, column=8, value="=" + GUARD.format(r=r, body=f'IFERROR(F{r}+G{r},"")'))
    wr.cell(row=r, column=16, value=f'=IFERROR(I{r}/H{r},"")')
    wr.cell(row=r, column=17, value=f'=IFERROR(K{r}/J{r},"")')
    wr.cell(row=r, column=18, value=f'=IFERROR(M{r}/L{r},"")')
    for cidx in range(3, 19):
        c = wr.cell(row=r, column=cidx)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        c.fill = auto_fill
        if cidx in (14, 15):
            c.number_format = '"$"#,##0'
        elif cidx in (16, 17, 18):
            c.number_format = "0.0%"
        else:
            c.number_format = "#,##0"

# ================================================================ KPI DASHBOARD
kd = wb.create_sheet("KPI Dashboard")
kd.sheet_view.showGridLines = False
banner(kd, "KPI DASHBOARD", 10)
subtitle(kd, "100% automatic — it reads your Daily Log. If a number looks wrong, the fix is in the Daily Log.", 10)
widths(kd, {"A": 32, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13, "G": 10,
            "H": 16, "I": 6, "J": 13})

periods = [("B", "This Week", "=TODAY()-WEEKDAY(TODAY(),2)+1", "=TODAY()"),
           ("C", "Last Week", "=B6-7", "=B6-1"),
           ("D", "This Month", "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "=TODAY()"),
           ("E", "Last Month", "=EOMONTH(TODAY(),-2)+1", "=EOMONTH(TODAY(),-1)"),
           ("F", "All Time", "=Setup!$C$8", "=TODAY()")]
for letter, name, s, e in periods:
    h = kd[f"{letter}4"]
    h.value = name
    h.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=PALETTE["accent"])
    h.alignment = Alignment(horizontal="center", vertical="center")
    lb = kd[f"{letter}5"]
    lb.value = f'=TEXT({letter}6,"d mmm")&" – "&TEXT({letter}7,"d mmm")'
    lb.font = Font(name=FONT, size=8, color=PALETTE["muted"])
    lb.alignment = Alignment(horizontal="center")
    for row, f in ((6, s), (7, e)):
        c = kd[f"{letter}{row}"]
        c.value = f
        c.number_format = "d mmm yyyy"
        c.font = Font(name=FONT, size=8, color=PALETTE["muted"])
        c.alignment = Alignment(horizontal="center")
kd.row_dimensions[4].height = 20
for row, txt in ((6, "Period start"), (7, "Period end")):
    c = label(kd, row, txt)
    c.font = Font(name=FONT, size=8, italic=True, color=PALETTE["muted"])


def metric_row(row, lbl, body_fn, fmt="#,##0", bold=False, indent=2):
    label(kd, row, lbl, bold=bold, indent=indent)
    for letter, *_ in periods:
        c = kd[f"{letter}{row}"]
        c.value = "=" + body_fn(letter)
        c.number_format = fmt
        c.font = Font(name=FONT, size=10, bold=bold)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        c.fill = auto_fill


section(kd, 9, "ACTIVITY — what you put in", 6)
metric_row(10, "Days Logged",
           lambda L: f'SUMIFS({LOGGED},{DATES},">="&{L}$6,{DATES},"<="&{L}$7)')
for row, lbl, key in ((11, "Reels / Posts", "reels"), (12, "Stories", "stories"),
                      (13, "Opener DMs Sent", "opener"), (14, "Follow-Up DMs Sent", "follow")):
    metric_row(row, lbl, lambda L, k=key: sumifs(M[k], f"{L}$6", f"{L}$7"))
metric_row(15, "Total DMs Sent", lambda L: f'{L}13+{L}14', bold=True)
metric_row(16, "Replies", lambda L: sumifs(M["replies"], f"{L}$6", f"{L}$7"))

section(kd, 18, "SALES — what it turned into", 6)
for row, lbl, key in ((19, "Calls Pitched", "pitched"), (20, "Calls Booked", "booked"),
                      (21, "Calls Showed", "showed"), (22, "Closes", "closes")):
    metric_row(row, lbl, lambda L, k=key: sumifs(M[k], f"{L}$6", f"{L}$7"))
metric_row(23, "Cash Collected", lambda L: sumifs(M["cash"], f"{L}$6", f"{L}$7"),
           fmt='"$"#,##0', bold=True)
metric_row(24, "Revenue Generated", lambda L: sumifs(M["revenue"], f"{L}$6", f"{L}$7"),
           fmt='"$"#,##0')

section(kd, 26, "CONVERSION RATES — where you're actually leaking", 6)
rates = [(27, "Reply Rate (replies ÷ DMs sent)", "16", "15", "0.0%"),
         (28, "Pitch Rate (pitched ÷ replies)", "19", "16", "0.0%"),
         (29, "Book Rate (booked ÷ pitched)", "20", "19", "0.0%"),
         (30, "Show Rate (showed ÷ booked)", "21", "20", "0.0%"),
         (31, "Close Rate (closes ÷ showed)", "22", "21", "0.0%"),
         (32, "DMs per Call Booked", "15", "20", "0.0"),
         (33, "Cash per Call Booked", "23", "20", '"$"#,##0')]
for row, lbl, num, den, fmt in rates:
    metric_row(row, lbl, lambda L, n=num, d=den: f'IFERROR({L}{n}/{L}{d},"")', fmt=fmt)
metric_row(34, "Cash per 100 DMs Sent",
           lambda L: f'IFERROR({L}23/{L}15*100,"")', fmt='"$"#,##0', bold=True)
kd["A34"].comment = Comment(
    "The number that makes the daily standards feel worth it: what one hundred DMs is "
    "actually worth to you in cash.", "The Called")

section(kd, 36, "DAILY STANDARDS — are you holding the line?", 6)
tg = kd["G36"]
tg.value = "Target"
tg.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
tg.fill = PatternFill("solid", fgColor=PALETTE["dark"])
tg.alignment = Alignment(horizontal="center")
for row, lbl, mrow, sref in ((37, "Avg Reels / day", "11", "C12"),
                             (38, "Avg Stories / day", "12", "C13"),
                             (39, "Avg Opener DMs / day", "13", "C14"),
                             (40, "Avg Follow-Up DMs / day", "14", "C15")):
    metric_row(row, lbl, lambda L, m=mrow: f'IFERROR({L}{m}/{L}10,"")', fmt="0.0")
    g = kd[f"G{row}"]
    g.value = f"=Setup!${sref[0]}${sref[1:]}"
    g.number_format = "0.0"
    g.font = Font(name=FONT, size=10, bold=True, color=PALETTE["accent"])
    g.alignment = Alignment(horizontal="center")
    g.fill = PatternFill("solid", fgColor=PALETTE["accent_soft"])
    g.border = BOX
metric_row(41, "Perfect Days (all 4 standards hit)",
           lambda L: (f'COUNTIFS({DATES},">="&{L}$6,{DATES},"<="&{L}$7,'
                      f'{col("O")},4)'))
metric_row(42, "Standards Hit Rate", lambda L: f'IFERROR({L}41/{L}10,"")',
           fmt="0.0%", bold=True)

kd.conditional_formatting.add("B37:F40", FormulaRule(
    formula=['AND(B37<>"",B37>=$G37)'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"]), stopIfTrue=False))
kd.conditional_formatting.add("B37:F40", FormulaRule(
    formula=['AND(B37<>"",B37<$G37)'], fill=red,
    font=Font(name=FONT, color=PALETTE["bad_text"]), stopIfTrue=False))

# month goal box
kd.merge_cells("H4:J4")
gh = kd["H4"]
gh.value = "THIS MONTH vs GOAL"
gh.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
gh.fill = PatternFill("solid", fgColor=PALETTE["accent"])
gh.alignment = Alignment(horizontal="center", vertical="center")
goal_rows = [(5, "Cash goal", "=Setup!$C$18", '"$"#,##0'),
             (6, "Collected so far", "=D23", '"$"#,##0'),
             (7, "% of goal", '=IFERROR(D23/Setup!$C$18,"")', "0%"),
             (8, "Days left in month", "=EOMONTH(TODAY(),0)-TODAY()", "#,##0")]
for row, lbl, f, fmt in goal_rows:
    kd.merge_cells(f"H{row}:I{row}")
    c = kd[f"H{row}"]
    c.value = lbl
    c.font = Font(name=FONT, size=9)
    c.alignment = Alignment(horizontal="left", indent=1)
    v = kd[f"J{row}"]
    v.value = f
    v.number_format = fmt
    v.font = Font(name=FONT, size=11, bold=True)
    v.alignment = Alignment(horizontal="center")
    v.fill = auto_fill
    v.border = BOX

section(kd, 44, "TRENDS — last 13 weeks", 10)


def styled_chart(chart, title, anchor):
    chart.title = title
    chart.height = 8
    chart.width = 15
    chart.style = 2
    kd.add_chart(chart, anchor)


cats = Reference(wr, min_col=2, min_row=WR_FIRST, max_row=WR_LAST)
c1 = LineChart()
c1.add_data(Reference(wr, min_col=8, max_col=9, min_row=3, max_row=WR_LAST), titles_from_data=True)
c1.set_categories(cats)
c1.y_axis.title = "Count"
styled_chart(c1, "DMs sent vs replies", "A45")

c2 = BarChart()
c2.type = "col"
c2.add_data(Reference(wr, min_col=11, min_row=3, max_row=WR_LAST), titles_from_data=True)
c2.add_data(Reference(wr, min_col=13, min_row=3, max_row=WR_LAST), titles_from_data=True)
c2.set_categories(cats)
styled_chart(c2, "Calls booked vs closes", "I45")

c3 = BarChart()
c3.type = "col"
c3.add_data(Reference(wr, min_col=14, min_row=3, max_row=WR_LAST), titles_from_data=True)
c3.set_categories(cats)
styled_chart(c3, "Cash collected per week", "A62")

c4 = LineChart()
c4.add_data(Reference(wr, min_col=16, max_col=17, min_row=3, max_row=WR_LAST), titles_from_data=True)
c4.set_categories(cats)
styled_chart(c4, "Reply rate vs book rate", "I62")

# ================================================================ INSTAGRAM TRACKER
ig = wb.create_sheet("Instagram Tracker")
ig.sheet_view.showGridLines = False
IG_FIRST = 20
IG_LAST = IG_FIRST + IG_WEEKS - 1
banner(ig, "INSTAGRAM TRACKER", 14)
subtitle(ig, "Weekly numbers off your Professional Dashboard, plus a monthly screenshot so your coach can audit the page.", 14)
widths(ig, {"A": 14, "B": 11, "C": 10, "D": 9, "E": 13, "F": 11, "G": 11, "H": 11,
            "I": 10, "J": 9, "K": 9, "L": 10, "M": 12, "N": 30})

section(ig, 4, "MONTHLY PROFILE CHECK — first Sunday of the month", 14)
steps = [
    "1.  Screenshot your profile — bio, follower count and the top of your grid.",
    "2.  Screenshot Professional Dashboard → Insights (last 30 days).",
    "3.  Click the box below, then Insert → Image → Image in cell, and upload.",
]
for i, s in enumerate(steps):
    c = ig.cell(row=5 + i, column=1, value=s)
    c.font = Font(name=FONT, size=10)
    ig.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=8)
    c.alignment = Alignment(horizontal="left", indent=1)

drop_border = Border(*(Side(style="medium", color=PALETTE["accent"]),) * 4)
for start_c, end_c, title in ((1, 5, "PROFILE SCREENSHOT"), (7, 11, "INSIGHTS SCREENSHOT")):
    t = ig.cell(row=9, column=start_c, value=title)
    t.font = Font(name=FONT, size=9, bold=True, color=PALETTE["accent"])
    ig.merge_cells(start_row=10, start_column=start_c, end_row=16, end_column=end_c)
    box = ig.cell(row=10, column=start_c, value="drop image here")
    box.font = Font(name=FONT, size=10, italic=True, color=PALETTE["muted"])
    box.fill = PatternFill("solid", fgColor=PALETTE["accent_soft"])
    box.alignment = Alignment(horizontal="center", vertical="center")
    for rr in range(10, 17):
        for cc in range(start_c, end_c + 1):
            ig.cell(row=rr, column=cc).border = drop_border

section(ig, 18, "WEEKLY NUMBERS — log every Sunday night", 14)
ig_headers = ["Week Starting", "Followers\n(end of week)", "Growth\n(+/-)", "Growth\n%",
              "Accounts\nReached", "Interactions", "Profile\nVisits", "New\nFollowers",
              "Link\nClicks", "Posts\n(auto)", "Stories\n(auto)", "Reach\nper Post",
              "New Followers\nper 1k Reached", "Notes"]
header_row(ig, 19, ig_headers)
ig.freeze_panes = "B20"

ig_inputs = list("BEFGHI") + ["N"]
for i in range(IG_WEEKS):
    r = IG_FIRST + i
    a = ig.cell(row=r, column=1, value="=Setup!$C$9" if i == 0 else f"=A{r-1}+7")
    a.number_format = "ddd d mmm"
    a.font = Font(name=FONT, size=10)
    a.fill = auto_fill
    a.border = BOX
    for letter in ig_inputs:
        c = ig[f"{letter}{r}"]
        c.fill = input_fill
        c.font = Font(name=FONT, size=10)
        c.border = BOX
        c.alignment = Alignment(horizontal="center" if letter != "N" else "left")
        if letter != "N":
            c.number_format = "#,##0"
    if i > 0:
        ig.cell(row=r, column=3, value=f'=IF(OR(B{r}="",B{r-1}=""),"",B{r}-B{r-1})')
        ig.cell(row=r, column=4, value=f'=IF(OR(B{r}="",B{r-1}=""),"",IFERROR((B{r}-B{r-1})/B{r-1},""))')
    for cidx, key in ((10, "reels"), (11, "stories")):
        body = f'SUMIFS({col(M[key])},{DATES},">="&$A{r},{DATES},"<="&$A{r}+6)'
        ig.cell(row=r, column=cidx, value=f'=IF($A{r}>TODAY(),"",{body})')
    ig.cell(row=r, column=12, value=f'=IFERROR(E{r}/J{r},"")')
    ig.cell(row=r, column=13, value=f'=IFERROR(H{r}/E{r}*1000,"")')
    for cidx, fmt in ((3, "+#,##0;-#,##0;0"), (4, "0.0%"), (10, "#,##0"), (11, "#,##0"),
                      (12, "0.0"), (13, "0.0")):
        c = ig.cell(row=r, column=cidx)
        c.number_format = fmt
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="center")
        c.fill = auto_fill
        c.border = BOX

dv_ig = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                       allow_blank=True, showErrorMessage=True,
                       errorTitle="Whole numbers only",
                       error="Enter a whole number of 0 or more.")
ig.add_data_validation(dv_ig)
for letter in "BEFGHI":
    dv_ig.add(f"{letter}{IG_FIRST}:{letter}{IG_LAST}")

ig.conditional_formatting.add(f"C{IG_FIRST}:D{IG_LAST}", FormulaRule(
    formula=[f'AND($C{IG_FIRST}<>"",$C{IG_FIRST}>0)'], fill=green,
    font=Font(name=FONT, bold=True, color=PALETTE["good_text"]), stopIfTrue=False))
ig.conditional_formatting.add(f"C{IG_FIRST}:D{IG_LAST}", FormulaRule(
    formula=[f'AND($C{IG_FIRST}<>"",$C{IG_FIRST}<0)'], fill=red,
    font=Font(name=FONT, bold=True, color=PALETTE["bad_text"]), stopIfTrue=False))

ig_cats = Reference(ig, min_col=1, min_row=IG_FIRST, max_row=IG_LAST)
ic1 = LineChart()
ic1.add_data(Reference(ig, min_col=2, min_row=19, max_row=IG_LAST), titles_from_data=True)
ic1.set_categories(ig_cats)
ic1.title = "Follower growth"
ic1.height, ic1.width, ic1.style = 8, 16, 2
ig.add_chart(ic1, "P4")

ic2 = BarChart()
ic2.type = "col"
ic2.add_data(Reference(ig, min_col=5, min_row=19, max_row=IG_LAST), titles_from_data=True)
ic2.set_categories(ig_cats)
ic2.title = "Accounts reached per week"
ic2.height, ic2.width, ic2.style = 8, 16, 2
ig.add_chart(ic2, "P22")

ig["J19"].comment = Comment(
    "Posts and Stories fill in automatically from your Daily Log — don't type here.",
    "The Called")

# ================================================================ finish
order = ["Start Here", "Setup", "Daily Log", "KPI Dashboard", "Instagram Tracker", "Weekly Rollup"]
wb._sheets = [wb[name] for name in order]
for name in order:
    wb[name].sheet_properties.tabColor = (
        PALETTE["accent"] if name in ("Daily Log", "Instagram Tracker", "Setup")
        else PALETTE["dark"])
wb.active = 0

DEMO = "--demo" in sys.argv
DEMO_DAYS = 45

if DEMO:
    rnd = random.Random(7)
    st["C5"] = "Demo Client"
    st["C6"] = "Francis"
    st["C7"] = "@democlient"
    st["C8"] = f"=TODAY()-{DEMO_DAYS}"

    followers = 3180
    prev_booked = 0
    for i in range(DEMO_DAYS):
        r = DL_FIRST + i
        light = i % 7 in (5, 6)          # two lighter days each week
        blank_day = i in (16, 31)        # days they went quiet but still logged
        if blank_day:
            vals = dict(reels=0, stories=0, opener=0, follow=0, replies=0,
                        pitched=0, booked=0, showed=0, closes=0, cash=0, revenue=0)
        elif light:
            vals = dict(reels=rnd.randint(2, 3), stories=rnd.randint(12, 22),
                        opener=rnd.randint(4, 8), follow=rnd.randint(3, 7),
                        replies=rnd.randint(1, 3), pitched=rnd.randint(0, 1),
                        booked=rnd.randint(0, 1), showed=0, closes=0,
                        cash=0, revenue=0)
        else:
            vals = dict(reels=rnd.randint(3, 5), stories=rnd.randint(28, 36),
                        opener=rnd.randint(9, 14), follow=rnd.randint(9, 13),
                        replies=rnd.randint(3, 7), pitched=rnd.randint(2, 4),
                        booked=rnd.randint(0, 2), showed=0, closes=0,
                        cash=0, revenue=0)
        # calls booked yesterday show up (or don't) today — gives a real show rate
        vals["showed"] = prev_booked if rnd.random() > 0.28 else max(0, prev_booked - 1)
        prev_booked = vals["booked"]
        if vals["showed"] > 0 and rnd.random() < 0.34:
            vals["closes"] = 1
            vals["cash"] = rnd.choice([1500, 2000, 3000])
            vals["revenue"] = vals["cash"] + rnd.choice([0, 0, 1500])
        followers += 0 if blank_day else rnd.randint(6, 28)
        for key, letter in M.items():
            if key == "followers":
                continue
            dl[f"{letter}{r}"] = vals[key]
        dl[f"N{r}"] = followers
    dl["P4"] = "First week — found my footing"
    dl[f"P{DL_FIRST + 16}"] = "Travel day, no work done"

    ig_rows = [(3180, 9800, 480, 310, 78, 22), (3268, 12400, 620, 395, 95, 31),
               (3351, 11100, 540, 350, 88, 26), (3452, 14600, 710, 430, 112, 38),
               (3530, 10250, 505, 322, 81, 24), (3641, 15800, 780, 468, 124, 41)]
    for i, (fol, reach, inter, visits, newf, clicks) in enumerate(ig_rows):
        r = IG_FIRST + i
        ig[f"B{r}"], ig[f"E{r}"], ig[f"F{r}"] = fol, reach, inter
        ig[f"G{r}"], ig[f"H{r}"], ig[f"I{r}"] = visits, newf, clicks
    ig[f"N{IG_FIRST + 4}"] = "Posted less — reach dropped with it"

out = ("docs/The_Called_Client_Tracker_DEMO.xlsx" if DEMO
       else "docs/The_Called_Client_Tracker.xlsx")
wb.save(out)
print(f"wrote {out}")
print(f"  Daily Log rows {DL_FIRST}-{DL_LAST} ({DAYS} days)")
print(f"  Weekly Rollup rows {WR_FIRST}-{WR_LAST} ({ROLL_WEEKS} rolling weeks)")
print(f"  Instagram rows {IG_FIRST}-{IG_LAST} ({IG_WEEKS} weeks)")
